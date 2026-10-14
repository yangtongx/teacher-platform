#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""analyze_exams.py - 考试成绩对比分析（完全离线，只吃本地 xlsx/csv）。

对应 workspaces/office/AGENTS.md 的「流程2：考试成绩分析」：

  读取两次成绩 -> 对比排名/总分/各科 -> 识别进退步学生 -> 生成分析报告

本脚本不做任何网络请求，也不调用任何大模型：学生姓名、学号、成绩全部只在
本机工作区内处理。报告里没有数据的项一律写「数据不足」，不编造结论。

用法示例：
  python3 skills/hwc-teacher-assistant/scripts/analyze_exams.py \
      --current data/grades/2026-04-10.xlsx \
      --previous data/grades/2026-03-10.xlsx \
      --output data/reports/2026-04-10-对比分析.md

  # 只看结构化结果
  python3 skills/hwc-teacher-assistant/scripts/analyze_exams.py \
      --current data/grades/2026-04-10.xlsx --previous data/grades/2026-03-10.xlsx --json

退出码：0 成功 / 2 输入不合法 / 3 缺少依赖。
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import statistics
import sys
from datetime import date as _date
from pathlib import Path

EXIT_OK = 0
EXIT_ARG = 2
EXIT_DEP = 3

# 无法用来算分的单元格内容，一律当缺考/缺失处理。
EMPTY_TOKENS = {
    "", "-", "--", "—", "－", "/", "\\", "?", "n/a", "na", "null", "none", "nan",
    "缺考", "缺", "无", "未考", "未参加", "免考", "弃考", "空", "不计分", "待录",
}

NAME_ALIASES = {
    "姓名", "学生姓名", "学生名字", "名字", "学生", "考生姓名", "name", "student", "studentname",
}
NAME_KEYS = ("姓名", "学生名字")
ID_ALIASES = {
    "学号", "学生学号", "学籍号", "考号", "准考证号", "座号", "编号", "考籍号",
    "id", "studentid", "studentno", "sno",
}
ID_KEYS = ("学号", "考号", "准考证号", "学籍号")
TOTAL_ALIASES = {
    "总分", "总成绩", "总分成绩", "总分合计", "总分数", "合计", "总计", "总和",
    "total", "totalscore", "sum",
}
TOTAL_KEYS = ("总分", "总成绩")

# 这些列是元信息，不当作科目来统计。
IGNORE_KEYS = (
    "排名", "名次", "班级", "年级", "学校", "性别", "备注", "序号", "考场",
    "考试", "日期", "等级", "等第", "编号", "平均", "家长", "电话", "联系", "班主任",
)
LONG_FORM_SUBJECT_KEYS = ("科目", "学科", "课程")
LONG_FORM_SCORE_KEYS = ("分数", "得分", "成绩")

# 满分推定候选：取观测最高分往上最接近的一个。
FULL_SCORE_CANDIDATES = (100, 120, 150, 200, 300)
PASS_RATIO_DEFAULT = 0.6
EXCELLENT_RATIO_DEFAULT = 0.85
WEAK_SUBJECT_RATE = 70.0


class SkillError(Exception):
    exit_code = EXIT_ARG


class InputError(SkillError):
    exit_code = EXIT_ARG


class MissingDependencyError(SkillError):
    exit_code = EXIT_DEP


# --------------------------------------------------------------------------
# 基础工具
# --------------------------------------------------------------------------


def clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


# 与 read_sheet.py 同口径的日期解析（脚本之间保持独立，不互相 import）。
_FULL_DATE_RE = re.compile(r"(20\d{2})\s*[-_./年]\s*(\d{1,2})\s*[-_./月]\s*(\d{1,2})\s*日?")
_COMPACT_DATE_RE = re.compile(r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)")
_YEAR_MONTH_RE = re.compile(r"(20\d{2})\s*[-_./年]\s*(\d{1,2})\s*月份?")


def parse_exam_date(text: str):
    """从文件名解析考试日期，返回 (日期字符串或 None, 说明)。"""
    raw = str(text or "")
    for pattern in (_FULL_DATE_RE, _COMPACT_DATE_RE):
        match = pattern.search(raw)
        if not match:
            continue
        year, month, day = (int(group) for group in match.groups())
        try:
            return _date(year, month, day).isoformat(), "来自文件名「" + match.group(0) + "」"
        except ValueError:
            return None, "文件名里的日期「" + match.group(0) + "」不是合法日期"
    match = _YEAR_MONTH_RE.search(raw)
    if match:
        year, month = (int(group) for group in match.groups())
        if 1 <= month <= 12:
            return _date(year, month, 1).isoformat(), "文件名只给到月份，按当月 1 日处理"
    return None, "文件名里没有可识别的日期"


def parse_score(value):
    """把单元格值转成分数（float）；空值/缺考/非数字一律返回 None。"""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        if number != number or number in (float("inf"), float("-inf")):
            return None
        return number
    text = str(value).strip()
    if not text:
        return None
    if text.lower().replace(" ", "") in EMPTY_TOKENS:
        return None
    candidate = text.rstrip("分").strip().replace(",", "").replace("，", "")
    if candidate.lower() in EMPTY_TOKENS:
        return None
    try:
        number = float(candidate)
    except ValueError:
        return None
    if number != number or number in (float("inf"), float("-inf")):
        return None
    return number


def normalize_header(value) -> str:
    """表头归一化：去空白、去星号、去括号注释、转小写，用于别名匹配。"""
    text = clean_text(value).replace("\u3000", "").replace(" ", "")
    text = text.strip("*＊")
    text = re.sub(r"[（(][^)）]*[)）]", "", text)
    return text.lower()


def infer_full_score(observed_max):
    """按观测最高分推定满分；超过候选项上限则返回 None（不瞎猜）。"""
    if observed_max is None or observed_max <= 0:
        return None
    for candidate in FULL_SCORE_CANDIDATES:
        if candidate >= observed_max - 1e-9:
            return candidate
    return None


def mean(values):
    return round(statistics.fmean(values), 2) if values else None


def percent(count, total):
    if not total:
        return None
    return round(count * 100.0 / total, 2)


def fmt(value, digits=2, suffix=""):
    """报告里统一格式化：None 一律写「数据不足」。"""
    if value is None:
        return "数据不足"
    if isinstance(value, float):
        text = f"{value:.{digits}f}".rstrip("0").rstrip(".")
        if text in ("", "-"):
            text = "0"
        return text + suffix
    return str(value) + suffix


def fmt_delta(value, digits=2, suffix=""):
    if value is None:
        return "数据不足"
    if value == 0:
        return "0" + suffix
    if isinstance(value, float):
        text = f"{value:+.{digits}f}".rstrip("0").rstrip(".")
    else:
        text = f"{value:+d}"
    return text + suffix


def fmt_rank_delta(value):
    """名次变化：正数表示名次前进（数字变小）。"""
    if value is None:
        return "数据不足"
    if value > 0:
        return "↑" + str(value)
    if value < 0:
        return "↓" + str(-value)
    return "持平"


# --------------------------------------------------------------------------
# 读取成绩表（xlsx / csv，唯一依赖是 openpyxl）
# --------------------------------------------------------------------------

HEADER_HINTS = ("姓名", "学号", "考号", "总分", "总成绩", "学生")


def find_header_row(rows):
    fallback = None
    for index, row in enumerate(rows[:10]):
        filled = [clean_text(cell) for cell in row if clean_text(cell)]
        if len(filled) < 2:
            continue
        if fallback is None:
            fallback = index
        if any(hint in cell for cell in filled for hint in HEADER_HINTS):
            return index
    return fallback


def table_from_rows(rows):
    header_index = find_header_row(rows)
    if header_index is None:
        return None, None
    header = [clean_text(cell) for cell in rows[header_index]]
    while header and not header[-1]:
        header.pop()
    if not header:
        return None, None

    data = []
    for row in rows[header_index + 1:]:
        cells = list(row)
        padded = [cells[i] if i < len(cells) else None for i in range(len(header))]
        if any(clean_text(cell) for cell in padded):
            data.append(padded)

    keep = [
        i for i in range(len(header))
        if header[i] or any(clean_text(row[i]) for row in data)
    ]
    if not keep:
        return None, None
    return [header[i] for i in keep], [[row[i] for i in keep] for row in data]


def read_xlsx_tables(path: Path, sheet_name=None):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - 部署环境已装
        raise MissingDependencyError("缺少依赖 openpyxl，无法读取 Excel 成绩文件。") from exc

    try:
        workbook = load_workbook(filename=str(path), data_only=True, read_only=True)
    except Exception as exc:
        raise InputError("无法读取 Excel 文件：" + str(path) + "\n原因：" + str(exc)) from exc

    names = workbook.sheetnames
    if sheet_name and sheet_name not in names:
        raise InputError(
            "文件里没有名为「" + str(sheet_name) + "」的工作表。实际工作表：" + "、".join(names)
        )
    targets = [sheet_name] if sheet_name else names
    tables = []
    for name in targets:
        rows = [list(row) for row in workbook[name].iter_rows(values_only=True)]
        header, data = table_from_rows(rows)
        if header:
            tables.append((name, header, data))
    workbook.close()
    if not tables:
        raise InputError("文件里没有读到可用的表头（前 10 行都没有表头行）：" + str(path))
    return tables


def read_csv_tables(path: Path):
    last_error = None
    for encoding in ("utf-8-sig", "gbk", "utf-8"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                rows = [list(row) for row in csv.reader(handle)]
        except UnicodeDecodeError as exc:
            last_error = exc
            continue
        header, data = table_from_rows(rows)
        if not header:
            raise InputError("CSV 里没有读到可用的表头：" + str(path))
        return [(path.stem, header, data)]
    raise InputError(
        "CSV 编码无法识别（试过 utf-8-sig / gbk / utf-8）：" + str(path) + "\n原因：" + str(last_error)
    )


def read_table(path: Path, sheet_name=None):
    """读成绩表，返回 (表头, 数据行, 提示信息列表, 使用的工作表名列表)。"""
    if not path.is_file():
        raise InputError("成绩文件不存在：" + str(path))
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        tables = read_xlsx_tables(path, sheet_name=sheet_name)
    elif suffix == ".csv":
        tables = read_csv_tables(path)
    elif suffix == ".xls":
        raise InputError(
            "openpyxl 不支持旧版 .xls 格式：" + str(path) + "\n请另存为 .xlsx 或 .csv 后重试。"
        )
    else:
        raise InputError("不支持的成绩文件类型：" + str(path) + "（只支持 .xlsx / .xlsm / .csv）")

    header, rows, notes, kept = None, [], [], []
    for name, table_header, table_rows in tables:
        if header is None:
            header, kept = list(table_header), [name]
            rows.extend(table_rows)
            continue
        if [clean_text(c) for c in table_header] == [clean_text(c) for c in header]:
            rows.extend(table_rows)
            kept.append(name)
        else:
            notes.append(
                "工作表「" + str(name) + "」的表头与「" + str(kept[0]) + "」不一致，本次未纳入统计；"
                "它的列名是：" + "、".join(clean_text(c) for c in table_header)
            )
    return header, rows, notes, kept


# --------------------------------------------------------------------------
# 列名识别（别名容错）
# --------------------------------------------------------------------------


def _find_column(normalized, aliases, keywords=()):
    for index, name in enumerate(normalized):
        if name and name in aliases:
            return index
    for index, name in enumerate(normalized):
        if name and any(keyword in name for keyword in keywords):
            return index
    return None


def detect_layout(header):
    """识别姓名/学号/总分/科目列；认不出来时抛中文错误并列出实际列名。"""
    normalized = [normalize_header(cell) for cell in header]

    for name in normalized:
        if name in LONG_FORM_SUBJECT_KEYS:
            for other in normalized:
                if other in LONG_FORM_SCORE_KEYS:
                    raise InputError(
                        "检测到长表格式（每行一科）：表里同时有「科目」和「成绩/分数」列。\n"
                        "本技能只处理宽表（一行一个学生、一列一科）。\n"
                        "请先在 Excel 里用数据透视表把它转成宽表，或另存一份宽表成绩。"
                    )

    name_index = _find_column(normalized, NAME_ALIASES, NAME_KEYS)
    if name_index is None:
        raise InputError(
            "无法识别学生姓名字段，不能继续分析。\n"
            "实际读到的列名：" + "、".join(clean_text(cell) for cell in header) + "\n"
            "支持的姓名别名：姓名 / 学生姓名 / 名字 / 学生 / name。\n"
            "请把表头改成上述名称之一后重试。"
        )

    id_index = _find_column(normalized, ID_ALIASES, ID_KEYS)
    total_index = _find_column(normalized, TOTAL_ALIASES, TOTAL_KEYS)
    if total_index == name_index:
        total_index = None
    if id_index == name_index:
        id_index = None

    subjects, ignored, non_numeric = [], [], []
    for index, cell in enumerate(header):
        if index in (name_index, id_index, total_index):
            continue
        key = normalized[index]
        if not key or any(ignore in key for ignore in IGNORE_KEYS):
            ignored.append(clean_text(cell))
            continue
        subjects.append((clean_text(cell), index))

    if not subjects:
        raise InputError(
            "没有识别到任何科目列，无法做成绩分析。\n"
            "实际读到的列名：" + "、".join(clean_text(cell) for cell in header) + "\n"
            "科目列 = 除姓名/学号/排名/班级等之外、且至少有一个分数的那一列。"
        )

    return {
        "name": name_index,
        "id": id_index,
        "total": total_index,
        "subjects": subjects,
        "ignored": [name for name in ignored if name],
        "header": [clean_text(cell) for cell in header],
    }


def _numeric_subjects(layout, rows):
    """留下真的有分数的科目列，其余归入「忽略（无有效分数）」。"""
    subjects, dropped = [], []
    for subject, index in layout["subjects"]:
        values = [parse_score(row[index]) for row in rows]
        if any(value is not None for value in values):
            subjects.append((subject, index))
        else:
            dropped.append(subject)
    return subjects, dropped


# --------------------------------------------------------------------------
# 单次考试成绩概况
# --------------------------------------------------------------------------


def load_exam(path, full_score=None, subject_full=None, pass_ratio=PASS_RATIO_DEFAULT,
              excellent_ratio=EXCELLENT_RATIO_DEFAULT, sheet_name=None):
    """读取一次考试，算出班级统计与名次。"""
    header, rows, notes, kept = read_table(Path(path), sheet_name=sheet_name)
    layout = detect_layout(header)
    subjects, dropped = _numeric_subjects(layout, rows)
    if not subjects:
        raise InputError(
            "所有科目列都没有有效分数：" + str(path) + "\n"
            "实际读到的列名：" + "、".join(layout["header"])
        )
    layout["subjects"] = subjects

    students = []
    for row in rows:
        name = clean_text(row[layout["name"]])
        if not name:
            continue
        student = {
            "name": name,
            "id": clean_text(row[layout["id"]]) if layout["id"] is not None else "",
            "subjects": {subject: parse_score(row[index]) for subject, index in subjects},
            "total": None,
            "total_source": None,
        }
        if layout["total"] is not None:
            student["total"] = parse_score(row[layout["total"]])
            if student["total"] is not None:
                student["total_source"] = "表格总分列"
        if student["total"] is None:
            available = [value for value in student["subjects"].values() if value is not None]
            if available:
                student["total"] = round(sum(available), 2)
                student["total_source"] = "各科求和"
        students.append(student)

    if not students:
        raise InputError("没有读到任何有姓名的学生行：" + str(path))

    rank_students(students)

    subject_full_info = {}
    manual = subject_full or {}
    for subject, index in subjects:
        values = [student["subjects"].get(subject) for student in students]
        values = [value for value in values if value is not None]
        if not values:
            continue
        if subject in manual:
            full, source = float(manual[subject]), "指定"
        else:
            full, source = infer_full_score(max(values)), "推定"
            if full is None:
                source = None
        subject_full_info[subject] = {"full": full, "source": source}

    if full_score is not None:
        full_total, full_total_source = float(full_score), "指定"
    else:
        known = [info["full"] for info in subject_full_info.values()]
        if known and all(value is not None for value in known) and len(known) == len(subjects):
            full_total, full_total_source = round(sum(known), 2), "推定"
        else:
            full_total, full_total_source = None, None

    exam = {
        "path": str(Path(path)),
        "date": None,
        "date_note": "",
        "sheets": kept,
        "notes": notes,
        "header": layout["header"],
        "ignored_columns": layout["ignored"],
        "dropped_columns": dropped,
        "subjects": [subject for subject, _ in subjects],
        "students": students,
        "full_total": full_total,
        "full_total_source": full_total_source,
        "pass_ratio": pass_ratio,
        "excellent_ratio": excellent_ratio,
        "subject_full": subject_full_info,
    }
    exam["stats"] = exam_stats(exam)
    return exam


def rank_students(students):
    """按总分降序排名，并列同名次（1,2,2,4）；无总分的名次为 None。"""
    ranked = sorted(
        [student for student in students if student["total"] is not None],
        key=lambda student: (-student["total"], student["name"]),
    )
    previous_total, previous_rank = None, 0
    for position, student in enumerate(ranked, 1):
        if previous_total is not None and student["total"] == previous_total:
            student["rank"] = previous_rank
        else:
            student["rank"] = position
            previous_rank, previous_total = position, student["total"]
    for student in students:
        student.setdefault("rank", None)
    return students


def exam_stats(exam):
    students = exam["students"]
    totals = [student["total"] for student in students if student["total"] is not None]
    full_total = exam["full_total"]
    stats = {
        "students": len(students),
        "with_total": len(totals),
        "missing_total": len(students) - len(totals),
        "avg": mean(totals),
        "max": round(max(totals), 2) if totals else None,
        "min": round(min(totals), 2) if totals else None,
        "full_total": full_total,
        "full_total_source": exam["full_total_source"],
        "pass_line": round(full_total * exam["pass_ratio"], 2) if full_total else None,
        "pass_count": None,
        "pass_rate": None,
        "excellent_line": round(full_total * exam["excellent_ratio"], 2) if full_total else None,
        "excellent_count": None,
        "excellent_rate": None,
        "subject_stats": {},
    }
    if full_total:
        stats["pass_count"] = sum(1 for value in totals if value >= stats["pass_line"])
        stats["pass_rate"] = percent(stats["pass_count"], len(totals))
        stats["excellent_count"] = sum(1 for value in totals if value >= stats["excellent_line"])
        stats["excellent_rate"] = percent(stats["excellent_count"], len(totals))

    for subject in exam["subjects"]:
        values = [
            student["subjects"].get(subject) for student in students
            if student["subjects"].get(subject) is not None
        ]
        info = exam["subject_full"].get(subject) or {}
        full = info.get("full")
        subject_stat = {
            "count": len(values),
            "missing": len(students) - len(values),
            "avg": mean(values),
            "max": round(max(values), 2) if values else None,
            "min": round(min(values), 2) if values else None,
            "full": full,
            "full_source": info.get("source"),
            "rate": None,
            "pass_rate": None,
            "excellent_rate": None,
        }
        if values and full:
            subject_stat["rate"] = round(statistics.fmean(values) * 100.0 / full, 2)
            subject_stat["pass_rate"] = percent(
                sum(1 for value in values if value >= full * exam["pass_ratio"]), len(values)
            )
            subject_stat["excellent_rate"] = percent(
                sum(1 for value in values if value >= full * exam["excellent_ratio"]), len(values)
            )
        stats["subject_stats"][subject] = subject_stat
    return stats


# --------------------------------------------------------------------------
# 两次考试对比
# --------------------------------------------------------------------------


def choose_matching_basis(current, previous):
    """两边大多数学号都齐全时按学号匹配，否则按姓名匹配。"""
    def ratio(exam):
        with_id = sum(1 for student in exam["students"] if student["id"])
        return with_id / len(exam["students"]) if exam["students"] else 0.0

    return "id" if ratio(current) >= 0.5 and ratio(previous) >= 0.5 else "name"


def compare_exams(current, previous, top=5):
    basis = choose_matching_basis(current, previous)
    previous_index = {}
    for student in previous["students"]:
        key = student["id"] if basis == "id" else student["name"]
        previous_index.setdefault(key, student)
    current_index = {}
    for student in current["students"]:
        key = student["id"] if basis == "id" else student["name"]
        current_index.setdefault(key, student)

    subjects = [name for name in current["subjects"]]
    for name in previous["subjects"]:
        if name not in subjects:
            subjects.append(name)

    matched, only_previous, only_current = [], [], []
    for key, before in previous_index.items():
        after = current_index.get(key)
        if after is None:
            only_previous.append(before["name"])
            continue
        total_delta = None
        if before["total"] is not None and after["total"] is not None:
            total_delta = round(after["total"] - before["total"], 2)
        rank_delta = None
        if before["rank"] is not None and after["rank"] is not None:
            rank_delta = before["rank"] - after["rank"]  # 正数 = 名次前进
        subject_delta = {}
        for subject in subjects:
            first = before["subjects"].get(subject)
            second = after["subjects"].get(subject)
            if first is not None and second is not None:
                subject_delta[subject] = round(second - first, 2)
        matched.append(
            {
                "name": after["name"],
                "id": after["id"] or before["id"],
                "previous_total": before["total"],
                "current_total": after["total"],
                "total_delta": total_delta,
                "previous_rank": before["rank"],
                "current_rank": after["rank"],
                "rank_delta": rank_delta,
                "subject_delta": subject_delta,
            }
        )
    for key, after in current_index.items():
        if key not in previous_index:
            only_current.append(after["name"])

    improved = sorted(
        [item for item in matched if item["total_delta"] is not None and item["total_delta"] > 0],
        key=lambda item: (-item["total_delta"], item["name"]),
    )[:top]
    declined = sorted(
        [item for item in matched if item["total_delta"] is not None and item["total_delta"] < 0],
        key=lambda item: (item["total_delta"], item["name"]),
    )[:top]
    rank_improved = sorted(
        [item for item in matched if item["rank_delta"] is not None and item["rank_delta"] > 0],
        key=lambda item: (-item["rank_delta"], item["name"]),
    )[:top]
    rank_declined = sorted(
        [item for item in matched if item["rank_delta"] is not None and item["rank_delta"] < 0],
        key=lambda item: (item["rank_delta"], item["name"]),
    )[:top]

    subject_diff = []
    for subject in subjects:
        before = (previous["stats"]["subject_stats"].get(subject) or {}).get("avg")
        after = (current["stats"]["subject_stats"].get(subject) or {}).get("avg")
        subject_diff.append(
            {
                "subject": subject,
                "previous_avg": before,
                "current_avg": after,
                "delta": round(after - before, 2) if before is not None and after is not None else None,
                "current_rate": (current["stats"]["subject_stats"].get(subject) or {}).get("rate"),
            }
        )
    subject_diff.sort(key=lambda item: (item["delta"] is None, item["delta"] if item["delta"] is not None else 0))

    weak_subjects = sorted(
        [
            {
                "subject": subject,
                "avg": stat["avg"],
                "rate": stat["rate"],
                "full": stat["full"],
                "full_source": stat["full_source"],
            }
            for subject, stat in current["stats"]["subject_stats"].items()
            if stat["rate"] is not None and stat["rate"] < WEAK_SUBJECT_RATE
        ],
        key=lambda item: (item["rate"], item["subject"]),
    )

    return {
        "basis": basis,
        "matched_count": len(matched),
        "only_previous": sorted(only_previous),
        "only_current": sorted(only_current),
        "students": sorted(
            matched,
            key=lambda item: (item["current_rank"] is None, item["current_rank"] or 0, item["name"]),
        ),
        "improved": improved,
        "declined": declined,
        "rank_improved": rank_improved,
        "rank_declined": rank_declined,
        "subject_diff": subject_diff,
        "weak_subjects": weak_subjects,
    }


def exam_trend(current, previous_list):
    exams = sorted(
        list(previous_list) + [current],
        key=lambda exam: (exam["date"] is None, exam["date"] or "", exam["path"]),
    )
    trend = []
    for exam in exams:
        stats = exam["stats"]
        trend.append(
            {
                "date": exam["date"],
                "label": Path(exam["path"]).name,
                "avg": stats["avg"],
                "max": stats["max"],
                "min": stats["min"],
                "pass_rate": stats["pass_rate"],
                "excellent_rate": stats["excellent_rate"],
            }
        )
    return trend


# --------------------------------------------------------------------------
# 报告生成（Markdown）
# --------------------------------------------------------------------------


def _table(headers, rows):
    lines = ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    for row in rows:
        lines.append("| " + " | ".join(str(cell) for cell in row) + " |")
    return "\n".join(lines)


def subject_heading(exam):
    parts = []
    for subject in exam["subjects"]:
        info = exam["subject_full"].get(subject) or {}
        if info.get("full"):
            parts.append(subject + " " + fmt(info["full"]) + "分(" + str(info.get("source") or "-") + ")")
        else:
            parts.append(subject + " 满分数据不足")
    return "、".join(parts)


def build_markdown(current, previous_list, comparisons, top):
    lines = []
    current_date = current["date"] or "未知日期"
    lines.append("# 考试成绩对比分析报告")
    lines.append("")
    lines.append("- 本次考试：" + current_date + "（" + Path(current["path"]).name + "）")
    if previous_list:
        for exam in previous_list:
            lines.append("- 对比考试：" + (exam["date"] or "未知日期") + "（" + Path(exam["path"]).name + "）")
    else:
        lines.append("- 对比考试：数据不足（未提供 --previous）")
    lines.append("- 本次样本：" + str(current["stats"]["students"]) + " 人，其中有总分 "
                 + str(current["stats"]["with_total"]) + " 人")
    lines.append("- 科目（本次）：" + "、".join(current["subjects"]))
    lines.append("- 满分口径：" + (
        "总分满分 " + fmt(current["full_total"]) + "分（" + str(current["full_total_source"]) + "）"
        if current["full_total"] else "数据不足（未提供 --full-score，且无法推定各科满分）"
    ))
    lines.append("- 及格线 / 优秀线：总分满分的 "
                 + fmt(current["pass_ratio"] * 100) + "% / " + fmt(current["excellent_ratio"] * 100) + "%")
    lines.append("")

    # 一、班级整体
    lines.append("## 一、班级整体")
    lines.append("")
    lines.append("### 1.1 本次考试概况")
    lines.append("")
    stats = current["stats"]
    lines.append(_table(
        ["指标", "数值"],
        [
            ["学生人数", stats["students"]],
            ["有总分人数", stats["with_total"]],
            ["缺总分人数", stats["missing_total"]],
            ["总分平均分", fmt(stats["avg"])],
            ["总分最高分", fmt(stats["max"])],
            ["总分最低分", fmt(stats["min"])],
            ["及格线", fmt(stats["pass_line"])],
            ["及格人数", fmt(stats["pass_count"])],
            ["及格率(%)", fmt(stats["pass_rate"])],
            ["优秀线", fmt(stats["excellent_line"])],
            ["优秀人数", fmt(stats["excellent_count"])],
            ["优秀率(%)", fmt(stats["excellent_rate"])],
        ],
    ))
    lines.append("")
    if stats["pass_rate"] is None:
        lines.append("> 及格率 / 优秀率：数据不足（总分满分未知，可用 --full-score 指定后重算）。")
        lines.append("")

    lines.append("### 1.2 各次考试趋势")
    lines.append("")
    trend = exam_trend(current, previous_list)
    if len(trend) <= 1:
        lines.append("数据不足（只有一次成绩，无法给出趋势）。")
    else:
        lines.append(_table(
            ["考试", "日期", "总分平均", "最高", "最低", "及格率(%)", "优秀率(%)"],
            [
                [
                    item["label"], item["date"] or "未知", fmt(item["avg"]), fmt(item["max"]),
                    fmt(item["min"]), fmt(item["pass_rate"]), fmt(item["excellent_rate"]),
                ]
                for item in trend
            ],
        ))
    lines.append("")

    # 二、各科情况
    lines.append("## 二、各科情况（本次）")
    lines.append("")
    lines.append("本次各科满分口径：" + subject_heading(current))
    lines.append("")
    lines.append(_table(
        ["科目", "有效人数", "缺失", "平均分", "最高", "最低", "满分", "得分率(%)", "及格率(%)", "优秀率(%)"],
        [
            [
                subject, stat["count"], stat["missing"], fmt(stat["avg"]), fmt(stat["max"]),
                fmt(stat["min"]),
                (fmt(stat["full"]) + "(" + str(stat["full_source"]) + ")") if stat["full"] else "数据不足",
                fmt(stat["rate"]), fmt(stat["pass_rate"]), fmt(stat["excellent_rate"]),
            ]
            for subject, stat in current["stats"]["subject_stats"].items()
        ],
    ))
    lines.append("")

    if comparisons:
        lines.append("### 2.1 各科平均分变化（相对基准考试）")
        lines.append("")
        for comparison in comparisons:
            lines.append("**对比 " + (comparison["previous"]["date"] or Path(comparison["previous"]["path"]).name)
                         + "（基准：" + ("是" if comparison["is_base"] else "否") + "）**")
            lines.append("")
            lines.append(_table(
                ["科目", "上次平均", "本次平均", "变化", "本次得分率(%)"],
                [
                    [
                        item["subject"], fmt(item["previous_avg"]), fmt(item["current_avg"]),
                        fmt_delta(item["delta"]), fmt(item["current_rate"]),
                    ]
                    for item in comparison["subject_diff"]
                ],
            ))
            lines.append("")

    lines.append("### " + ("2.2" if comparisons else "2.1") + " 单科薄弱点提示")
    lines.append("")
    base = comparisons[-1] if comparisons else None
    weak = base["weak_subjects"] if base else [
        {
            "subject": subject, "avg": stat["avg"], "rate": stat["rate"],
            "full": stat["full"], "full_source": stat["full_source"],
        }
        for subject, stat in sorted(
            current["stats"]["subject_stats"].items(),
            key=lambda item: (item[1]["avg"] is None, item[1]["avg"] or 0),
        )
        if stat["rate"] is not None and stat["rate"] < WEAK_SUBJECT_RATE
    ]
    if weak:
        for item in weak:
            lines.append(
                "- " + item["subject"] + "：平均分 " + fmt(item["avg"]) + "，得分率 "
                + fmt(item["rate"], suffix="%") + "（低于 " + fmt(WEAK_SUBJECT_RATE, suffix="%")
                + "，满分按" + str(item.get("full_source") or "-") + " " + fmt(item["full"]) + " 分计）"
            )
    elif any(stat["rate"] is not None for stat in current["stats"]["subject_stats"].values()):
        lines.append("- 各科得分率均不低于 " + fmt(WEAK_SUBJECT_RATE, suffix="%") + "，本次没有明显的单科薄弱点。")
    else:
        lines.append("- 数据不足（无法确定各科满分，得分率算不出来；可用 --subject-full 语文=150 指定）。")
    lines.append("")

    # 三、名次变化
    lines.append("## 三、学生总分与名次变化")
    lines.append("")
    if not base:
        lines.append("数据不足（未提供对比成绩 --previous，无法计算名次升降）。")
        lines.append("")
    else:
        lines.append("基准考试：" + (base["previous"]["date"] or Path(base["previous"]["path"]).name)
                     + "；匹配方式：" + ("按学号" if base["basis"] == "id" else "按姓名")
                     + "；匹配到 " + str(base["matched_count"]) + " 人。")
        lines.append("")
        lines.append("> 名次变化：↑ 表示名次前进几位，↓ 表示后退几位（班级第一名最好）。")
        lines.append("")
        lines.append(_table(
            ["姓名", "上次总分", "本次总分", "总分变化", "上次名次", "本次名次", "名次变化"],
            [
                [
                    item["name"], fmt(item["previous_total"]), fmt(item["current_total"]),
                    fmt_delta(item["total_delta"]), fmt(item["previous_rank"]),
                    fmt(item["current_rank"]), fmt_rank_delta(item["rank_delta"]),
                ]
                for item in base["students"]
            ],
        ))
        lines.append("")
        if base["only_previous"]:
            lines.append("- 只在对比考试里出现的学生：" + "、".join(base["only_previous"]))
        if base["only_current"]:
            lines.append("- 只在本次考试里出现的学生：" + "、".join(base["only_current"]))
        if base["only_previous"] or base["only_current"]:
            lines.append("")

    # 四、进步学生
    lines.append("## 四、进步学生（按总分变化，前 " + str(top) + " 名）")
    lines.append("")
    if not base:
        lines.append("数据不足（未提供对比成绩 --previous）。")
    elif not base["improved"]:
        lines.append("本次没有总分较上次提高的学生。")
    else:
        lines.append(_table(
            ["姓名", "上次总分", "本次总分", "总分变化", "名次变化"],
            [
                [item["name"], fmt(item["previous_total"]), fmt(item["current_total"]),
                 fmt_delta(item["total_delta"]), fmt_rank_delta(item["rank_delta"])]
                for item in base["improved"]
            ],
        ))
    lines.append("")

    # 五、退步学生
    lines.append("## 五、退步学生（按总分变化，前 " + str(top) + " 名）")
    lines.append("")
    if not base:
        lines.append("数据不足（未提供对比成绩 --previous）。")
    elif not base["declined"]:
        lines.append("本次没有总分较上次下降的学生。")
    else:
        lines.append(_table(
            ["姓名", "上次总分", "本次总分", "总分变化", "名次变化"],
            [
                [item["name"], fmt(item["previous_total"]), fmt(item["current_total"]),
                 fmt_delta(item["total_delta"]), fmt_rank_delta(item["rank_delta"])]
                for item in base["declined"]
            ],
        ))
    lines.append("")

    # 六、名次升降（独立于总分变化）
    lines.append("## 六、名次进步 / 退步名单（前 " + str(top) + " 名）")
    lines.append("")
    if not base:
        lines.append("数据不足（未提供对比成绩 --previous）。")
    else:
        if base["rank_improved"]:
            lines.append("**名次进步**")
            lines.append("")
            lines.append(_table(
                ["姓名", "上次名次", "本次名次", "名次变化", "总分变化"],
                [
                    [item["name"], fmt(item["previous_rank"]), fmt(item["current_rank"]),
                     fmt_rank_delta(item["rank_delta"]), fmt_delta(item["total_delta"])]
                    for item in base["rank_improved"]
                ],
            ))
        else:
            lines.append("名次进步：本次没有名次前进的学生。")
        lines.append("")
        if base["rank_declined"]:
            lines.append("**名次退步**")
            lines.append("")
            lines.append(_table(
                ["姓名", "上次名次", "本次名次", "名次变化", "总分变化"],
                [
                    [item["name"], fmt(item["previous_rank"]), fmt(item["current_rank"]),
                     fmt_rank_delta(item["rank_delta"]), fmt_delta(item["total_delta"])]
                    for item in base["rank_declined"]
                ],
            ))
        else:
            lines.append("名次退步：本次没有名次后退的学生。")
    lines.append("")

    # 七、教学建议（只根据上面这些数字）
    lines.append("## 七、教学建议（仅基于本次报告里的数据）")
    lines.append("")
    suggestions = build_suggestions(current, base, weak)
    if suggestions:
        for item in suggestions:
            lines.append("- " + item)
    else:
        lines.append("- 数据不足，本次不给建议。")
    lines.append("")

    # 八、数据说明
    lines.append("## 八、数据说明与缺口")
    lines.append("")
    lines.append("- 本次表头：" + "、".join(current["header"]))
    if current["ignored_columns"]:
        lines.append("- 本次未作为科目的列：" + "、".join(current["ignored_columns"]))
    if current["dropped_columns"]:
        lines.append("- 本次无有效分数、已忽略的列：" + "、".join(current["dropped_columns"]))
    if current["notes"]:
        for note in current["notes"]:
            lines.append("- " + note)
    if current["stats"]["missing_total"]:
        lines.append("- 本次有 " + str(current["stats"]["missing_total"])
                       + " 人没有总分（缺考或该列为空），未纳入平均分与名次计算。")
    for subject, stat in current["stats"]["subject_stats"].items():
        if stat["missing"]:
            lines.append("- " + subject + " 有 " + str(stat["missing"]) + " 人没有有效分数。")
    if previous_list:
        for exam in previous_list:
            if exam["notes"]:
                for note in exam["notes"]:
                    lines.append("- (" + Path(exam["path"]).name + ") " + note)
    lines.append("- 满分口径为推定值时，仅按本次观测最高分推断，可能与实际满分不同；"
                 "如需精确的及格率/优秀率，请用 --full-score 与 --subject-full 指定。")
    lines.append("")
    lines.append("> 本报告由教师办公助手在本机离线生成，学生数据未离开工作区，未调用任何外部服务。")
    return "\n".join(lines)


def build_suggestions(current, base, weak):
    """只把已经算出来的数字翻译成建议，不添加任何数据之外的判断。"""
    suggestions = []
    if weak:
        subject = weak[0]["subject"]
        suggestions.append(
            subject + "是本次得分率最低的科目（得分率 " + fmt(weak[0]["rate"], suffix="%")
            + "），建议安排一次针对 " + subject + " 的专项讲评与练习。"
        )
    if base:
        for item in base["subject_diff"]:
            if item["delta"] is not None and item["delta"] <= -3:
                suggestions.append(
                    item["subject"] + "平均分较上次下降 " + fmt(abs(item["delta"]))
                    + " 分，建议先核对这次试卷难度与教学进度，再决定是否补课。"
                )
        if base["declined"]:
            names = "、".join(item["name"] for item in base["declined"][:3])
            suggestions.append(
                "总分退步最明显的 " + str(len(base["declined"])) + " 人中有：" + names
                + "，建议逐一沟通了解原因。"
            )
        if base["improved"]:
            names = "、".join(item["name"] for item in base["improved"][:3])
            suggestions.append("总分进步的学生（如 " + names + "）建议在班上做一次经验分享。")
        if current["stats"]["avg"] is not None and base["previous"]["stats"]["avg"] is not None:
            gap = round(current["stats"]["avg"] - base["previous"]["stats"]["avg"], 2)
            if abs(gap) < 1:
                suggestions.append("班级总分平均分与上次基本持平（变化 " + fmt_delta(gap) + " 分）。")
    if current["stats"]["pass_rate"] is not None and current["stats"]["pass_rate"] < 60:
        suggestions.append(
            "本次总分及格率 " + fmt(current["stats"]["pass_rate"], suffix="%")
            + "，低于 60%，建议评估是否需要分层辅导。"
        )
    return suggestions


def build_json(current, previous_list, comparisons, top):
    def exam_public(exam):
        return {
            "path": exam["path"],
            "date": exam["date"],
            "date_note": exam["date_note"],
            "sheets": exam["sheets"],
            "header": exam["header"],
            "subjects": exam["subjects"],
            "ignored_columns": exam["ignored_columns"],
            "dropped_columns": exam["dropped_columns"],
            "full_total": exam["full_total"],
            "full_total_source": exam["full_total_source"],
            "pass_ratio": exam["pass_ratio"],
            "excellent_ratio": exam["excellent_ratio"],
            "subject_full": exam["subject_full"],
            "stats": exam["stats"],
            "notes": exam["notes"],
            "students": exam["students"],
        }

    return {
        "current": exam_public(current),
        "previous": [exam_public(exam) for exam in previous_list],
        "trend": exam_trend(current, previous_list),
        "comparisons": [
            {
                "is_base": comparison["is_base"],
                "previous_path": comparison["previous"]["path"],
                "previous_date": comparison["previous"]["date"],
                "basis": comparison["basis"],
                "matched_count": comparison["matched_count"],
                "only_previous": comparison["only_previous"],
                "only_current": comparison["only_current"],
                "students": comparison["students"],
                "improved": comparison["improved"],
                "declined": comparison["declined"],
                "rank_improved": comparison["rank_improved"],
                "rank_declined": comparison["rank_declined"],
                "subject_diff": comparison["subject_diff"],
                "weak_subjects": comparison["weak_subjects"],
            }
            for comparison in comparisons
        ],
        "top": top,
        "weak_subject_rate_threshold": WEAK_SUBJECT_RATE,
        "note": "本结果由 analyze_exams.py 离线生成，学生数据未离开工作区。",
    }


# --------------------------------------------------------------------------
# 命令行
# --------------------------------------------------------------------------


def parse_subject_full(items):
    mapping = {}
    for item in items or []:
        if "=" not in item:
            raise InputError("--subject-full 需要写成 科目=满分，例如 --subject-full 语文=150；收到：" + str(item))
        name, value = item.split("=", 1)
        name = name.strip()
        try:
            full = float(value.strip())
        except ValueError:
            raise InputError("--subject-full 的满分不是数字：" + str(item)) from None
        if not name or full <= 0:
            raise InputError("--subject-full 参数不合法：" + str(item))
        mapping[name] = full
    return mapping


def build_parser():
    parser = argparse.ArgumentParser(
        description="考试成绩对比分析：班级统计、名次升降、进退步名单（离线，读本地 xlsx/csv）",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("--current", required=True, help="本次成绩文件（.xlsx / .csv）")
    parser.add_argument("--previous", action="append", default=[],
                        help="对比成绩文件，可多次传入；名次升降以日期最近的一次为基准")
    parser.add_argument("--output", help="Markdown 报告输出路径（默认 data/reports/{文件名}-对比分析.md；"
                                        "以 .json 结尾时写结构化 JSON）")
    parser.add_argument("--top", type=int, default=5, help="进步/退步名单条数，默认 5")
    parser.add_argument("--json", action="store_true", help="把结构化结果打印到 stdout")
    parser.add_argument("--full-score", type=float, help="总分满分（用于及格率/优秀率；默认按观测值推定）")
    parser.add_argument("--subject-full", action="append", default=[],
                        help="指定某科满分，可多次传入，如 --subject-full 语文=150")
    parser.add_argument("--pass-ratio", type=float, default=PASS_RATIO_DEFAULT,
                        help="及格线占总分满分的比例，默认 0.6")
    parser.add_argument("--excellent-ratio", type=float, default=EXCELLENT_RATIO_DEFAULT,
                        help="优秀线占总分满分的比例，默认 0.85")
    parser.add_argument("--sheet", help="只分析指定工作表（默认合并表头一致的全部工作表）")
    parser.add_argument("--data-dir", default="data", help="数据目录，默认 data（相对工作区根目录）")
    return parser


def main(argv=None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    try:
        if args.top < 1:
            raise InputError("--top 必须大于等于 1，收到：" + str(args.top))
        if not 0 < args.pass_ratio <= 1:
            raise InputError("--pass-ratio 必须落在 (0, 1] 区间，收到：" + str(args.pass_ratio))
        if not 0 < args.excellent_ratio <= 1:
            raise InputError("--excellent-ratio 必须落在 (0, 1] 区间，收到：" + str(args.excellent_ratio))
        subject_full = parse_subject_full(args.subject_full)
    except SkillError as exc:
        print("错误：" + str(exc), file=sys.stderr)
        return exc.exit_code

    try:
        current = load_exam(args.current, full_score=args.full_score, subject_full=subject_full,
                            pass_ratio=args.pass_ratio, excellent_ratio=args.excellent_ratio,
                            sheet_name=args.sheet)
        previous_list = []
        for path in args.previous:
            previous_list.append(
                load_exam(path, full_score=args.full_score, subject_full=subject_full,
                          pass_ratio=args.pass_ratio, excellent_ratio=args.excellent_ratio,
                          sheet_name=args.sheet)
            )
    except SkillError as exc:
        print("错误：" + str(exc), file=sys.stderr)
        return exc.exit_code
    except OSError as exc:
        print("错误：文件读写失败：" + str(exc), file=sys.stderr)
        return EXIT_ARG

    for exam in [current] + previous_list:
        exam["date"], exam["date_note"] = parse_exam_date(Path(exam["path"]).name)

    previous_list.sort(key=lambda exam: (exam["date"] or "", exam["path"]))
    comparisons = []
    for index, exam in enumerate(previous_list):
        comparison = compare_exams(current, exam, top=args.top)
        comparison["previous"] = exam
        comparison["is_base"] = index == len(previous_list) - 1
        comparisons.append(comparison)

    report = build_markdown(current, previous_list, comparisons, args.top)
    payload = build_json(current, previous_list, comparisons, args.top)

    if args.output:
        output = Path(args.output)
    else:
        output = Path(args.data_dir) / "reports" / (Path(args.current).stem + "-对比分析.md")

    if output.suffix.lower() == ".json":
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        written = output
    else:
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(report, encoding="utf-8")
        written = output

    # --json 时 stdout 必须只有这一个 JSON 文档（可直接管道给 jq / json.load），
    # 人类可读的摘要一律走 stderr。
    summary = sys.stderr if args.json else sys.stdout

    if args.json:
        print(json.dumps(payload, ensure_ascii=False, indent=2))

    print("[完成] 分析报告已写入：" + str(written.resolve()), file=summary)
    print("       本次考试：" + str(current["date"]) + "（" + Path(current["path"]).name + "）",
          file=summary)
    if previous_list:
        print("       对比考试：" + "、".join(
            (exam["date"] or "未知") for exam in previous_list) + "；基准="
            + str(comparisons[-1]["previous"]["date"]), file=summary)
        print("       匹配方式：" + ("按学号" if comparisons[-1]["basis"] == "id" else "按姓名")
              + "，匹配到 " + str(comparisons[-1]["matched_count"]) + " 人", file=summary)
        print("       进步 " + str(len(comparisons[-1]["improved"])) + " 人 / 退步 "
              + str(len(comparisons[-1]["declined"])) + " 人（报告里各列前 " + str(args.top) + " 名）",
              file=summary)
    else:
        print("       对比考试：无（未提供 --previous，报告里相关小节标注为数据不足）", file=summary)
    print("       总分平均分：" + fmt(current["stats"]["avg"])
          + "；及格率：" + fmt(current["stats"]["pass_rate"], suffix="%")
          + "；优秀率：" + fmt(current["stats"]["excellent_rate"], suffix="%"), file=summary)
    print("       提示：本脚本完全离线，学生数据未发送给任何外部服务。", file=summary)
    return EXIT_OK


if __name__ == "__main__":
    sys.exit(main())
