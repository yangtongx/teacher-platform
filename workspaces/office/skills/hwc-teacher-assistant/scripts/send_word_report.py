#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""send_word_report.py - 生成学生成绩 Word 报告，并（可选）一对一私发给家长。

对应 workspaces/office/AGENTS.md 的「流程3：发送成绩报告」与「流程4：私聊通知」：

  获取最近 N 次成绩 -> 生成 Word 报告 -> 私聊/邮件附带发送

两条红线在本脚本里是硬约束：
  1. 默认只生成、不发送（dry-run）。要真正发送必须显式加 --send。
  2. 只接受家长的个人 open_id（ou_ 开头）。传入群 chat_id（oc_ 开头）一律拒绝，
     因为家校沟通规则是「私聊 = 一对一私发，绝不在群里发」。

用法示例：
  # 只生成（默认，不发消息）
  python3 skills/hwc-teacher-assistant/scripts/send_word_report.py \
      --student 张三 --grades data/grades/2026-03-10.xlsx data/grades/2026-04-10.xlsx \
      --output data/reports/张三-成绩报告.docx

  # 生成并一对一私发
  python3 skills/hwc-teacher-assistant/scripts/send_word_report.py \
      --student 张三 --to-open-id ou_xxxx --send

  # 不给 --grades 时，自动取 data/grades/ 下最近 3 次
  python3 skills/hwc-teacher-assistant/scripts/send_word_report.py --student 张三 --send --to-open-id ou_xxxx

退出码：0 成功 / 2 输入不合法 / 3 缺少凭据或依赖 / 4 飞书接口或网络失败 / 5 触发安全红线（拒绝群发）。
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from pathlib import Path

# 同目录的 analyze_exams.py：列名识别与统计口径必须与流程2完全一致。
sys.path.insert(0, str(Path(__file__).resolve().parent))

try:
    import analyze_exams as analysis
except ImportError as _import_error:  # 只把本脚本单独拷走时会走到这里
    analysis = None
    _ANALYSIS_IMPORT_ERROR = _import_error

BASE_URL = os.environ.get("FEISHU_BASE_URL", "https://open.feishu.cn").rstrip("/")
TIMEOUT = 15
DEFAULT_LIMIT = 3
FONT_NAME = "宋体"

EXIT_OK = 0
EXIT_ARG = 2
EXIT_CRED = 3
EXIT_API = 4
EXIT_FORBIDDEN = 5

# 群聊/部门/union_id 等非个人 open_id 的前缀，命中即拒绝发送。
NON_PERSONAL_PREFIXES = {
    "oc_": "群聊 chat_id",
    "chat_": "群聊 chat_id",
    "od_": "部门 ID",
    "on_": "union_id",
}


class SkillError(Exception):
    exit_code = EXIT_ARG


class InputError(SkillError):
    exit_code = EXIT_ARG


class MissingCredentialError(SkillError):
    exit_code = EXIT_CRED


class MissingDependencyError(SkillError):
    exit_code = EXIT_CRED


class FeishuAPIError(SkillError):
    exit_code = EXIT_API


class GroupSendForbidden(SkillError):
    """命中「绝不在群里发」这条红线。"""

    exit_code = EXIT_FORBIDDEN


def mask_id(value: str) -> str:
    text = str(value or "")
    if len(text) <= 8:
        return "*" * len(text)
    return text[:4] + "*" * (len(text) - 8) + text[-4:]


# --------------------------------------------------------------------------
# 依赖与凭据（与 read_sheet.py 同口径：环境变量优先，其次 ~/.openclaw/.env）
# --------------------------------------------------------------------------


def _import_docx():
    try:
        from docx import Document  # noqa: F401
    except ImportError as exc:
        raise MissingDependencyError(
            "缺少依赖 python-docx，无法生成 Word 报告。服务器上应已安装 python-docx 1.2.0。"
        ) from exc


def parse_env_file(path) -> dict:
    values = {}
    try:
        content = Path(path).read_text(encoding="utf-8", errors="replace")
    except OSError:
        return values
    for line in content.splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        if line.startswith("export "):
            line = line[len("export "):].strip()
        if "=" not in line:
            continue
        key, value = line.split("=", 1)
        key, value = key.strip(), value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in "\"'":
            value = value[1:-1]
        if key:
            values[key] = value
    return values


def credentials_path() -> Path:
    override = (os.environ.get("FEISHU_ENV_FILE") or "").strip()
    if override:
        return Path(override).expanduser()
    return Path.home() / ".openclaw" / ".env"


def diagnose_missing_credentials(missing) -> str:
    path = credentials_path()
    return (
        "飞书凭据未配置：缺少 " + "、".join(missing) + "\n"
        "本脚本只有加 --send 时才需要凭据；只生成报告不需要。\n"
        "配置方式（凭据只放服务器上，绝不要写进代码或提交仓库）：\n"
        "  1) 环境变量：export FEISHU_APP_ID=cli_xxx; export FEISHU_APP_SECRET=xxx\n"
        "  2) 凭据文件：" + str(path) + "\n"
        "       FEISHU_APP_ID=cli_xxx\n"
        "       FEISHU_APP_SECRET=xxx\n"
        "     权限建议收紧：chmod 600 " + str(path) + "\n"
        "所需权限：im:message（发消息）、im:chat.members:read（读群成员）、sheets:spreadsheet:readonly（读成绩表）。"
    )


def load_credentials():
    app_id = (os.environ.get("FEISHU_APP_ID") or "").strip()
    app_secret = (os.environ.get("FEISHU_APP_SECRET") or "").strip()
    source = "环境变量 FEISHU_APP_ID / FEISHU_APP_SECRET"
    if not (app_id and app_secret):
        path = credentials_path()
        if path.is_file():
            values = parse_env_file(path)
            app_id = app_id or (values.get("FEISHU_APP_ID") or "").strip()
            app_secret = app_secret or (values.get("FEISHU_APP_SECRET") or "").strip()
            source = "凭据文件 " + str(path)
    missing = [
        name
        for name, value in (("FEISHU_APP_ID", app_id), ("FEISHU_APP_SECRET", app_secret))
        if not value
    ]
    if missing:
        raise MissingCredentialError(diagnose_missing_credentials(missing))
    return app_id, app_secret, source


# --------------------------------------------------------------------------
# HTTP 封装（独立函数，便于测试 monkeypatch；统一超时）
# --------------------------------------------------------------------------


def _decode_json(response, url: str) -> dict:
    try:
        body = response.json()
    except ValueError as exc:
        raise FeishuAPIError(
            "飞书接口返回的不是 JSON：HTTP " + str(getattr(response, "status_code", "?")) + " " + url
            + "\n响应片段：" + str(getattr(response, "text", ""))[:300]
        ) from exc
    if isinstance(body, dict):
        code = body.get("code")
        if code not in (0, None):
            message = body.get("msg") or body.get("message") or "(无 msg)"
            raise FeishuAPIError(
                "飞书接口报错 code=" + str(code) + " msg=" + str(message) + "\n请求：" + url
            )
    if getattr(response, "status_code", 200) >= 400:
        raise FeishuAPIError(
            "飞书接口返回 HTTP " + str(response.status_code) + "：" + url
            + "\n响应片段：" + str(getattr(response, "text", ""))[:300]
        )
    return body


def http_post_json(url: str, payload: dict, timeout: int = TIMEOUT, headers=None, params=None) -> dict:
    import requests

    try:
        response = requests.post(
            url, json=payload, params=params, timeout=timeout, headers=headers or {}
        )
    except requests.RequestException as exc:
        raise FeishuAPIError("请求飞书接口失败（网络错误）：" + url + "\n原因：" + str(exc)) from exc
    return _decode_json(response, url)


def http_post_multipart(url: str, headers, files, data=None, timeout: int = TIMEOUT) -> dict:
    import requests

    try:
        response = requests.post(
            url, headers=headers or {}, files=files, data=data or {}, timeout=timeout
        )
    except requests.RequestException as exc:
        raise FeishuAPIError("上传文件到飞书失败（网络错误）：" + url + "\n原因：" + str(exc)) from exc
    return _decode_json(response, url)


def get_tenant_access_token(app_id: str, app_secret: str, timeout: int = TIMEOUT):
    url = BASE_URL + "/open-apis/auth/v3/tenant_access_token/internal"
    body = http_post_json(url, {"app_id": app_id, "app_secret": app_secret}, timeout=timeout)
    token = body.get("tenant_access_token")
    if not token:
        raise FeishuAPIError(
            "没有取到 tenant_access_token：code=" + str(body.get("code")) + " msg="
            + str(body.get("msg")) + "\n常见原因：app_id / app_secret 写错，或应用未启用。"
        )
    return token, int(body.get("expire") or 0)


# --------------------------------------------------------------------------
# 发送（强制一对一）
# --------------------------------------------------------------------------


def assert_one_to_one(receive_id: str) -> str:
    """只放行个人 open_id；群 chat_id 等一律拒绝（红线：绝不在群里发）。"""
    value = str(receive_id or "").strip()
    if not value:
        raise InputError("--to-open-id 为空，无法确定收件人。")
    for prefix, label in NON_PERSONAL_PREFIXES.items():
        if value.startswith(prefix):
            raise GroupSendForbidden(
                "拒绝发送：--to-open-id 收到的是" + label + "（" + mask_id(value) + "），不是家长的个人 open_id。\n"
                "家校沟通红线：私聊必须一对一私发，绝不在群里发。\n"
                "请改用家长个人的 open_id（ou_ 开头）。"
            )
    if not value.startswith("ou_"):
        raise InputError(
            "--to-open-id 只接受家长个人 open_id（ou_ 开头），收到：" + mask_id(value) + "\n"
            "如果不确定，请先用群成员列表接口按姓名匹配出 open_id 再发送。"
        )
    return value


def upload_file(access_token: str, path: Path, timeout: int = TIMEOUT) -> str:
    """上传 Word 文件，返回 file_key。"""
    url = BASE_URL + "/open-apis/im/v1/files"
    with Path(path).open("rb") as handle:
        files = {"file": (Path(path).name, handle, "application/octet-stream")}
        body = http_post_multipart(
            url,
            headers={"Authorization": "Bearer " + access_token},
            files=files,
            data={"file_type": "stream", "file_name": Path(path).name},
            timeout=timeout,
        )
    file_key = (body.get("data") or {}).get("file_key")
    if not file_key:
        raise FeishuAPIError("上传文件后没有拿到 file_key，飞书返回：" + json.dumps(body, ensure_ascii=False)[:300])
    return file_key


def send_text_message(access_token: str, open_id: str, text: str, timeout: int = TIMEOUT) -> dict:
    """一对一发给个人 open_id（receive_id_type 固定为 open_id）。"""
    receive_id = assert_one_to_one(open_id)
    url = BASE_URL + "/open-apis/im/v1/messages"
    body = http_post_json(
        url,
        {"receive_id": receive_id, "msg_type": "text", "content": json.dumps({"text": text}, ensure_ascii=False)},
        timeout=timeout,
        headers={"Authorization": "Bearer " + access_token},
        params={"receive_id_type": "open_id"},
    )
    return body.get("data") or {}


def send_file_message(access_token: str, open_id: str, file_key: str, timeout: int = TIMEOUT) -> dict:
    """把报告作为文件消息一对一私发。"""
    receive_id = assert_one_to_one(open_id)
    url = BASE_URL + "/open-apis/im/v1/messages"
    body = http_post_json(
        url,
        {"receive_id": receive_id, "msg_type": "file", "content": json.dumps({"file_key": file_key})},
        timeout=timeout,
        headers={"Authorization": "Bearer " + access_token},
        params={"receive_id_type": "open_id"},
    )
    return body.get("data") or {}


# --------------------------------------------------------------------------
# 取「最近 N 次成绩」
# --------------------------------------------------------------------------


def list_recent_grades(data_dir: Path, limit: int, explicit=None):
    """返回按日期升序排列的成绩文件列表（最多 limit 个）。"""
    if explicit:
        paths = [Path(item) for item in explicit]
        for path in paths:
            if not path.is_file():
                raise InputError("成绩文件不存在：" + str(path))
    else:
        grades_dir = Path(data_dir) / "grades"
        if not grades_dir.is_dir():
            raise InputError(
                "找不到成绩目录：" + str(grades_dir) + "\n"
                "请先用 read_sheet.py 导入成绩（--url 或 --local），或显式传 --grades。"
            )
        paths = sorted(
            [item for item in grades_dir.iterdir() if item.suffix.lower() in (".xlsx", ".xlsm", ".csv")],
            key=lambda item: item.name,
        )
        if not paths:
            raise InputError(
                "成绩目录里没有文件：" + str(grades_dir) + "\n"
                "请先用 read_sheet.py 导入成绩，或显式传 --grades。"
            )
    if limit and limit > 0:
        paths = sorted(paths, key=lambda item: item.name)[-limit:]
    return paths


def pick_student(exam, name, student_id=None):
    if student_id:
        matches = [student for student in exam["students"] if student["id"] == str(student_id)]
        if not matches:
            return None
        return matches[0]
    matches = [student for student in exam["students"] if student["name"] == name]
    if len(matches) > 1:
        raise InputError(
            "在 " + exam["path"] + " 里找到多个叫「" + name + "」的学生，"
            "请用 --student-id 指定学号后重试。"
        )
    return matches[0] if matches else None


# --------------------------------------------------------------------------
# Word 报告
# --------------------------------------------------------------------------


def _set_style_font(style, name=FONT_NAME, size=None):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    style.font.name = name
    if size is not None:
        style.font.size = size
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attribute), name)


def set_run_font(run, name=FONT_NAME, size=None, bold=None):
    """中文必须显式设置 w:eastAsia，否则 Word 里会退化成默认西文字体。"""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    run.font.name = name
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attribute), name)


def add_paragraph(document, text, size=None, bold=False, align=None):
    from docx.shared import Pt

    paragraph = document.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size or Pt(10.5), bold=bold)
    return paragraph


def add_table(document, headers, rows):
    from docx.shared import Pt

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(text))
        set_run_font(run, size=Pt(10.5), bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row):
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run("" if text is None else str(text))
            set_run_font(run, size=Pt(10.5))
    return table


def build_report(student_name, student_id, exams, student_rows, warnings, output: Path,
                 class_name=None, teacher_note=None):
    """生成 Word 成绩报告。

    exams: 按日期升序的考试列表（已 load_exam）
    student_rows: 与 exams 一一对应的学生记录（可能为 None 表示该次无该生成绩）
    """
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    _import_docx()
    document = Document()
    _set_style_font(document.styles["Normal"], size=Pt(10.5))

    add_paragraph(document, student_name + " 成绩报告", size=Pt(18), bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)

    meta = "姓名：" + student_name
    if student_id:
        meta += "    学号：" + str(student_id)
    if class_name:
        meta += "    班级：" + str(class_name)
    add_paragraph(document, meta)
    if exams:
        span = (exams[0]["date"] or Path(exams[0]["path"]).stem) + " 至 " + (
            exams[-1]["date"] or Path(exams[-1]["path"]).stem)
        add_paragraph(document, "报告范围：" + span + "（共 " + str(len(exams)) + " 次考试）")

    # 表1：总分与名次
    add_paragraph(document, "一、总成绩与班级名次", size=Pt(12), bold=True)
    rows = []
    previous_rank = None
    for exam, student in zip(exams, student_rows):
        date = exam["date"] or Path(exam["path"]).stem
        if student is None:
            rows.append([date, "未找到该学生成绩", "-", "-", "-"])
            previous_rank = None
            continue
        total_text = analysis.fmt(student["total"])
        rank = student.get("rank")
        rank_text = str(rank) if rank is not None else "数据不足"
        if rank is not None and previous_rank is not None:
            delta = previous_rank - rank
            change = "↑" + str(delta) if delta > 0 else ("↓" + str(-delta) if delta < 0 else "持平")
        else:
            change = "—"
        rows.append([date, total_text, analysis.fmt(exam["stats"]["avg"]), rank_text, change])
        previous_rank = rank
    add_table(document, ["考试日期", "总分", "班级平均分", "班级排名", "名次变化"], rows)

    # 表2：各科成绩
    subjects = []
    for exam in exams:
        for subject in exam["subjects"]:
            if subject not in subjects:
                subjects.append(subject)
    add_paragraph(document, "二、各科成绩", size=Pt(12), bold=True)
    headers = ["科目"] + [(exam["date"] or Path(exam["path"]).stem) for exam in exams]
    if len(exams) >= 2:
        headers.append("变化")
    subject_rows = []
    for subject in subjects:
        values = []
        for exam, student in zip(exams, student_rows):
            if student is None:
                values.append(None)
            else:
                values.append(student["subjects"].get(subject))
        row = [subject] + [analysis.fmt(value) for value in values]
        if len(exams) >= 2:
            if values[-1] is not None and values[-2] is not None:
                row.append(analysis.fmt_delta(round(values[-1] - values[-2], 2)))
            else:
                row.append("数据不足")
        subject_rows.append(row)
    total_row = ["总分"] + [
        analysis.fmt(student["total"] if student else None) for student in student_rows
    ]
    if len(exams) >= 2:
        totals = [student["total"] if student else None for student in student_rows]
        if totals[-1] is not None and totals[-2] is not None:
            total_row.append(analysis.fmt_delta(round(totals[-1] - totals[-2], 2)))
        else:
            total_row.append("数据不足")
    subject_rows.append(total_row)
    add_table(document, headers, subject_rows)

    # 表3：最近一次的各科明细
    latest_exam, latest_student = exams[-1], student_rows[-1]
    if latest_student is not None:
        add_paragraph(document, "三、最近一次考试各科明细（"
                      + (latest_exam["date"] or Path(latest_exam["path"]).stem) + "）",
                      size=Pt(12), bold=True)
        detail_rows = []
        for subject in latest_exam["subjects"]:
            score = latest_student["subjects"].get(subject)
            stat = latest_exam["stats"]["subject_stats"].get(subject) or {}
            detail_rows.append([
                subject,
                analysis.fmt(score),
                analysis.fmt(stat.get("avg")),
                "数据不足" if stat.get("full") is None else analysis.fmt(stat["full"]) + "（" + str(stat.get("full_source")) + "）",
                analysis.fmt(stat.get("rate"), suffix="%"),
            ])
        add_table(document, ["科目", "个人得分", "班级平均分", "满分", "班级得分率"], detail_rows)

    # 说明
    add_paragraph(document, "四、数据来源与说明", size=Pt(12), bold=True)
    for exam in exams:
        add_paragraph(document, "- " + (exam["date"] or "日期未知") + "：" + Path(exam["path"]).name
                      + "（班级 " + str(exam["stats"]["students"]) + " 人；总分平均分 "
                      + analysis.fmt(exam["stats"]["avg"]) + "）")
    full_total = exams[-1]["full_total"] if exams else None
    add_paragraph(document, "- 满分口径："
                  + ("总分满分 " + analysis.fmt(full_total) + "（" + str(exams[-1]["full_total_source"]) + "）"
                     if full_total else "数据不足（未指定 --full-score，无法给出及格率/优秀率）"))
    for warning in warnings:
        add_paragraph(document, "- " + warning)
    if teacher_note:
        add_paragraph(document, "- " + teacher_note)
    add_paragraph(document, "本报告由教师办公助手在本机生成，成绩数据只在工作区内处理，"
                            "未发送给任何第三方服务。请家长本人查收，不要转发到群聊。", size=Pt(9))

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    return output


# --------------------------------------------------------------------------
# 命令行
# --------------------------------------------------------------------------


def build_parser():
    parser = argparse.ArgumentParser(
        description="生成学生成绩 Word 报告；默认只生成不发送，加 --send 才一对一私发",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("--student", required=True, help="学生姓名")
    parser.add_argument("--student-id", help="学号（同名学生必填，用于区分）")
    parser.add_argument("--grades", nargs="+", help="成绩文件，可多个；不传则取 data/grades/ 下最近 --limit 次")
    parser.add_argument("--output", help="输出 docx 路径（默认 data/reports/{姓名}-成绩报告.docx）")
    parser.add_argument("--to-open-id", help="家长个人 open_id（ou_ 开头）；群 chat_id（oc_ 开头）会被拒绝")
    parser.add_argument("--send", action="store_true", help="真正发送（默认只生成不发送）")
    parser.add_argument("--message", help="随报告一起发的文本内容，默认用内置模板")
    parser.add_argument("--class-name", help="班级名称，写进报告抬头（可选）")
    parser.add_argument("--limit", type=int, default=DEFAULT_LIMIT, help="合并最近几次成绩，默认 3")
    parser.add_argument("--data-dir", default="data", help="数据目录，默认 data（相对工作区根目录）")
    parser.add_argument("--timeout", type=int, default=TIMEOUT, help="HTTP 超时秒数，默认 15")
    return parser


def sanitize_filename(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", str(name)).strip("_") or "student"


def main(argv=None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    access_token = None
    try:
        _import_docx()
        if args.limit < 1:
            raise InputError("--limit 必须大于等于 1，收到：" + str(args.limit))

        # 凭据检查放在最前面：--send 时缺凭据要立刻失败，不能先假装成功。
        if args.send:
            if not args.to_open_id:
                raise InputError(
                    "--send 需要同时提供 --to-open-id（家长个人 open_id，ou_ 开头）。\n"
                    "红线：私聊只做一对一私发，不接受群 ID。"
                )
            # 在取凭据之前就拦住群 ID：红线判定不依赖任何网络请求。
            assert_one_to_one(args.to_open_id)
            app_id, app_secret, cred_source = load_credentials()
            access_token, _ = get_tenant_access_token(app_id, app_secret, timeout=args.timeout)

        if analysis is None:
            raise MissingDependencyError(
                "找不到同目录的 analyze_exams.py（列名识别与统计口径由它提供）："
                + str(_ANALYSIS_IMPORT_ERROR)
            )

        paths = list_recent_grades(Path(args.data_dir), args.limit, explicit=args.grades)

        loaded, student_rows, warnings = [], [], []
        for path in paths:
            exam = analysis.load_exam(path)
            exam["date"], exam["date_note"] = analysis.parse_exam_date(Path(exam["path"]).name)
            student = pick_student(exam, args.student, student_id=args.student_id)
            if student is None:
                warnings.append(
                    "该次考试没有找到「" + args.student + "」的成绩（缺考或未录入），本次不列入报告："
                    + Path(exam["path"]).name
                )
                continue
            loaded.append(exam)
            student_rows.append(student)

        if not loaded:
            latest_names = []
            try:
                probe = analysis.load_exam(paths[-1])
                latest_names = [item["name"] for item in probe["students"]][:10]
            except SkillError:
                latest_names = []
            hint = ("最近一次成绩（" + Path(paths[-1]).name + "）里的学生有："
                    + "、".join(latest_names) + "\n") if latest_names else ""
            raise InputError(
                "在选定的 " + str(len(paths)) + " 个成绩文件里都没有找到学生「" + args.student + "」。\n"
                + hint + "请核对姓名（要跟成绩表里的写法完全一致），或用 --student-id 指定学号。"
            )

        order = sorted(range(len(loaded)), key=lambda index: (loaded[index]["date"] or "", loaded[index]["path"]))
        loaded = [loaded[index] for index in order]
        student_rows = [student_rows[index] for index in order]

        student_id = args.student_id or student_rows[-1]["id"]
        output = Path(args.output) if args.output else (
            Path(args.data_dir) / "reports" / (sanitize_filename(args.student) + "-成绩报告.docx")
        )
        build_report(args.student, student_id, loaded, student_rows, warnings, output,
                     class_name=args.class_name)
        if not output.is_file() or output.stat().st_size <= 0:
            raise InputError("Word 报告没有生成成功：" + str(output))

        print("[完成] Word 报告已生成：" + str(output.resolve()))
        print("       学生：" + args.student + (("（学号 " + str(student_id) + "）") if student_id else ""))
        print("       包含考试：" + "、".join(
            (exam["date"] or Path(exam["path"]).stem) for exam in loaded) + "（共 " + str(len(loaded)) + " 次）")
        for warning in warnings:
            print("       [提示] " + warning)

        if not args.send:
            print("")
            print("[演练模式] 未发送任何消息（默认只生成不发送）。")
            print("       收件人：未指定（要发送请加 --send --to-open-id ou_xxx）")
            print("       附件：" + str(output.resolve()))
            print("       红线提醒：私聊只做一对一私发，绝不在群里发。")
            return EXIT_OK

        print("")
        print("[发送计划]")
        print("       收件人（家长 open_id）：" + args.to_open_id)
        print("       学生：" + args.student + (("（学号 " + str(student_id) + "）") if student_id else ""))
        print("       附件：" + str(output.resolve()))
        print("       发送方式：飞书 IM 一对一私聊（receive_id_type=open_id），不会发到任何群里")

        text = args.message or (
            args.student + "家长您好，附上孩子最近 " + str(len(loaded)) + " 次考试的成绩报告（"
            + "、".join((exam["date"] or Path(exam["path"]).stem) for exam in loaded)
            + "），请查收。本报告仅您本人可见，如有疑问欢迎随时沟通。"
        )
        send_text_message(access_token, args.to_open_id, text, timeout=args.timeout)
        file_key = upload_file(access_token, output, timeout=args.timeout)
        result = send_file_message(access_token, args.to_open_id, file_key, timeout=args.timeout)
        print("[完成] 已一对一私发文本消息 + Word 报告附件。")
        print("       message_id：" + str(result.get("message_id") or "(飞书未返回)"))
        print("       学生数据只在工作区内处理，未发送给任何第三方服务（飞书 IM 除外，且仅发给该家长本人）。")
        return EXIT_OK

    except SkillError as exc:
        print("错误：" + str(exc), file=sys.stderr)
        return exc.exit_code
    except OSError as exc:
        print("错误：文件读写失败：" + str(exc), file=sys.stderr)
        return EXIT_API


if __name__ == "__main__":
    sys.exit(main())
