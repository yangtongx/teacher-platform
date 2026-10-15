import io
import os
import tempfile
import unittest
import zipfile

if not os.environ.get("DATA_DIR"):
    os.environ["DATA_DIR"] = tempfile.mkdtemp()
os.environ.setdefault("ADMIN_USERNAME", "teacher")
os.environ.setdefault("ADMIN_PASSWORD", "test-password-1")
os.environ.setdefault("ADMIN_DISPLAY_NAME", "测试老师")

from app import db
from app.auth import hash_password
from app.extract import extract_text


class HardeningTests(unittest.TestCase):
    def setUp(self):
        db.init_db()

    def tearDown(self):
        with db.connect() as conn:
            conn.execute("DELETE FROM rate_events")
            conn.commit()

    def test_docs_are_closed_and_responses_are_framed(self):
        from fastapi.testclient import TestClient

        from app.main import app

        client = TestClient(app)
        self.assertEqual(client.get("/docs").status_code, 404)
        self.assertEqual(client.get("/openapi.json").status_code, 404)
        health = client.get("/healthz")
        self.assertEqual(health.headers["x-frame-options"], "DENY")
        self.assertIn("frame-ancestors 'none'", health.headers["content-security-policy"])
        self.assertEqual(health.headers["x-content-type-options"], "nosniff")

    def test_null_origin_on_same_site_logout_reaches_login(self):
        from fastapi.testclient import TestClient

        from app.main import app

        client = TestClient(app)
        response = client.post(
            "/logout",
            headers={"Origin": "null", "Sec-Fetch-Site": "same-origin"},
            follow_redirects=False,
        )
        self.assertEqual(response.status_code, 303)
        self.assertEqual(response.headers["location"], "/login")
        self.assertEqual(
            client.get("/healthz").headers["referrer-policy"],
            "strict-origin-when-cross-origin",
        )

    def test_cross_origin_post_is_rejected(self):
        from fastapi.testclient import TestClient

        from app.main import app

        client = TestClient(app)
        response = client.post(
            "/login",
            data={"username": "teacher", "password": "test-password-1"},
            headers={"Origin": "https://evil.example"},
            follow_redirects=False,
        )
        self.assertEqual(response.status_code, 403)

    def test_null_origin_from_another_site_is_rejected(self):
        from fastapi.testclient import TestClient

        from app.main import app

        client = TestClient(app)
        response = client.post(
            "/logout",
            headers={"Origin": "null", "Sec-Fetch-Site": "cross-site"},
            follow_redirects=False,
        )
        self.assertEqual(response.status_code, 403)

    def test_password_change_drops_the_session(self):
        user = db.user_by_username("teacher")
        token = db.create_session(user["id"], "lesson")
        self.assertIsNotNone(db.session_user(token))
        db.set_password_hash(user["id"], hash_password("another-password"))
        self.assertIsNone(db.session_user(token))
        db.set_password_hash(user["id"], hash_password("test-password-1"))

    def test_login_lockout_uses_the_same_message_channel(self):
        from fastapi.testclient import TestClient

        from app.main import app

        os.environ["LOGIN_FAIL_LIMIT"] = "2"
        try:
            client = TestClient(app)
            for _ in range(2):
                failed = client.post(
                    "/login",
                    data={"username": "teacher", "password": "wrong-password"},
                    follow_redirects=False,
                )
                self.assertEqual(failed.status_code, 303)
                self.assertIn("error=1", failed.headers["location"])
            locked = client.post(
                "/login",
                data={"username": "teacher", "password": "test-password-1"},
                follow_redirects=False,
            )
            self.assertEqual(locked.status_code, 303)
            self.assertIn("error=2", locked.headers["location"])
        finally:
            os.environ.pop("LOGIN_FAIL_LIMIT", None)

    def test_download_expires(self):
        from fastapi.testclient import TestClient

        from app.main import app

        user = db.user_by_username("teacher")
        db.add_message(user["id"], "office", "assistant", "一份通知")
        message = db.list_messages(user["id"], "office")[-1]
        db.save_document(user["id"], message["id"], "expired-token", "通知.docx")
        db.save_document(user["id"], message["id"], "fresh-token", "通知.docx")
        docs = db.DATA_DIR / "docs"
        docs.mkdir(parents=True, exist_ok=True)
        (docs / "expired-token.docx").write_bytes(b"expired")
        (docs / "fresh-token.docx").write_bytes(b"fresh")
        with db.connect() as conn:
            conn.execute(
                "UPDATE documents SET expires_at = ? WHERE token = ?",
                ("2000-01-01T00:00:00+00:00", "expired-token"),
            )
            conn.commit()
        client = TestClient(app)
        logged = client.post(
            "/login",
            data={"username": "teacher", "password": "test-password-1"},
            follow_redirects=False,
        )
        self.assertEqual(logged.status_code, 303)
        self.assertEqual(client.get("/files/fresh-token").status_code, 200)
        self.assertEqual(client.get("/files/expired-token").status_code, 404)

    def test_oversized_docx_is_rejected(self):
        from app import extract

        previous = extract.MAX_UNCOMPRESSED
        extract.MAX_UNCOMPRESSED = 50
        try:
            buffer = io.BytesIO()
            with zipfile.ZipFile(buffer, "w") as archive:
                archive.writestr("word/document.xml", b"a" * 80)
            with self.assertRaises(ValueError):
                extract_text("big.docx", buffer.getvalue())
        finally:
            extract.MAX_UNCOMPRESSED = previous

    def test_scanned_pdf_is_recognized_and_text_pdf_is_not(self):
        from io import BytesIO

        from pypdf import PdfWriter

        from app import extract

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.add_blank_page(width=200, height=200)
        buffer = BytesIO()
        writer.write(buffer)
        scanned = buffer.getvalue()
        seen = []

        def fake_ocr(_data, index, _holder):
            seen.append(index)
            return "勾股定理" if index == 0 else ""

        previous_ocr = extract._ocr_page
        previous_limit = extract.MAX_OCR_PAGES
        extract._ocr_page = fake_ocr
        extract.MAX_OCR_PAGES = 1
        try:
            text = extract.extract_text("scan.pdf", scanned)
            self.assertEqual(seen, [0])
            self.assertTrue(text.startswith(extract.OCR_NOTE))
            self.assertIn("勾股定理", text)
            self.assertIn("只识别了前 1 页", text)
            extract._ocr_page = lambda *_args: ""
            with self.assertRaises(ValueError) as empty:
                extract.extract_text("scan.pdf", scanned)
            self.assertIn("没有识别出文字", str(empty.exception))
        finally:
            extract._ocr_page = previous_ocr
            extract.MAX_OCR_PAGES = previous_limit

        seen.clear()

        def reject_ocr(*_args):
            seen.append("called")
            return "不该识别"

        def text_pdf(message: str) -> bytes:
            stream = f"BT /F1 18 Tf 36 120 Td ({message}) Tj ET\n".encode()
            objects = [
                b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n",
                b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n",
                b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 400 200] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n",
                b"4 0 obj\n<< /Length %d >>\nstream\n" % len(stream) + stream + b"endstream\nendobj\n",
                b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n",
            ]
            out = bytearray(b"%PDF-1.4\n")
            offsets = [0]
            for obj in objects:
                offsets.append(len(out))
                out.extend(obj)
            xref = len(out)
            out.extend(b"xref\n0 6\n0000000000 65535 f \n")
            for offset in offsets[1:]:
                out.extend(f"{offset:010d} 00000 n \n".encode())
            out.extend(f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
            return bytes(out)

        extract._ocr_page = reject_ocr
        try:
            text = extract.extract_text(
                "lesson.pdf",
                text_pdf("Third unit reviews fractions before the lesson starts."),
            )
            self.assertIn("fractions", text)
            self.assertNotIn(extract.OCR_NOTE, text)
            self.assertEqual(seen, [])
        finally:
            extract._ocr_page = previous_ocr

    def test_image_blocks_keep_column_order_and_repair_english(self):
        from app.image_ocr import needs_english_repair, order_blocks

        blocks = [
            {"x0": 420, "y0": 40, "x1": 700, "y1": 80, "kind": "title"},
            {"x0": 40, "y0": 120, "x1": 300, "y1": 180, "kind": "text"},
            {"x0": 420, "y0": 120, "x1": 700, "y1": 180, "kind": "text"},
            {"x0": 40, "y0": 40, "x1": 300, "y1": 80, "kind": "title"},
            {"x0": 40, "y0": 200, "x1": 200, "y1": 280, "kind": "figure"},
        ]
        ordered = order_blocks(blocks, 760)
        self.assertEqual([(item["x0"], item["y0"]) for item in ordered], [(40, 40), (40, 120), (420, 40), (420, 120)])
        self.assertFalse(needs_english_repair("勾股定理的证明", 0.4))
        self.assertTrue(needs_english_repair("fractions before class", 0.4))
        self.assertFalse(needs_english_repair("fractions before class", 0.95))

    def test_ppt_rejects_unsupported_knowledge_and_downloads(self):
        from io import BytesIO

        from fastapi.testclient import TestClient
        from pptx import Presentation

        from app.main import app
        from app.ppt import review_slides, write_ppt

        rejected = review_slides(
            [
                {"role": "课题与目标", "title": "分数", "bullets": ["认识几分之一"], "source": "教案"},
                {"role": "新知", "title": "假分数", "bullets": ["分子大于分母"], "source": "要求"},
                {"role": "小结", "title": "记住", "bullets": ["分数表示一份"], "source": "教案"},
                {"role": "作业", "title": "练习", "bullets": ["做课本练习"], "source": "课本"},
            ]
        )
        self.assertTrue(any("新知" in item for item in rejected))

        slides = [
            {"role": "课题与目标", "title": "认识分数", "bullets": ["能说出几分之一"], "source": "教案", "notes": "先看图"},
            {"role": "导入", "title": "分月饼", "bullets": ["一块月饼平均分成四份"], "source": "要求"},
            {"role": "新知", "title": "几分之一", "bullets": ["份数是分母"], "source": "课本"},
            {"role": "小结", "title": "今天记住", "bullets": ["分母表示平均分成的份数"], "source": "教案"},
            {"role": "作业", "title": "课后", "bullets": ["完成练习二"], "source": "课本"},
        ]
        db.init_db()
        user = db.user_by_username("teacher")
        conversation_id = db.conversation_id(user["id"], "lesson")
        db.add_message(user["id"], "lesson", "user", "做课件")
        saved = write_ppt(user, conversation_id, "认识分数", slides)
        self.assertTrue(saved["name"].endswith("_课件.pptx"))
        client = TestClient(app)
        logged = client.post(
            "/login",
            data={"username": "teacher", "password": "test-password-1"},
            follow_redirects=False,
        )
        self.assertEqual(logged.status_code, 303)
        downloaded = client.get(saved["url"])
        self.assertEqual(downloaded.status_code, 200)
        presentation = Presentation(BytesIO(downloaded.content))
        self.assertEqual(len(presentation.slides), 5)
        cover = "\n".join(
            shape.text_frame.text
            for shape in presentation.slides[0].shapes
            if getattr(shape, "has_text_frame", False)
        )
        self.assertIn("认识分数", cover)
        self.assertNotIn("Your lesson title", cover)
        self.assertNotIn("Learning objective", cover)
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        self.assertTrue(
            any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in presentation.slides[0].shapes)
        )

    def test_ppt_style_stays_on_the_account_and_drops_old_text(self):
        from io import BytesIO

        from pptx import Presentation
        from pptx.util import Inches, Pt

        from app.ppt import write_ppt
        from app.ppt_style import StyleNeeded, describe_style, prepare_export, store_pptx

        db.init_db()
        user = db.user_by_username("teacher")
        conversation_id = db.conversation_id(int(user["id"]), "main")
        db.add_message(int(user["id"]), "main", "user", "做课件")
        self.assertIn("直接用", describe_style(user, conversation_id))
        self.assertNotIn("用 AI 自己的", describe_style(user, conversation_id))
        template, subject = prepare_export(user, conversation_id, "subject", "数学")
        self.assertEqual(subject, "数学")
        self.assertIn("slate", template.name)
        self.assertIn("请确认", describe_style(user, conversation_id))
        with self.assertRaises(StyleNeeded) as caught:
            prepare_export(user, conversation_id, "", "")
        self.assertIn("仍用你上次的课件样式", str(caught.exception))
        kept, _name = prepare_export(user, conversation_id, "keep", "")
        self.assertEqual(kept, template)
        self.assertIn("请确认", describe_style(user, 0))

        with db.connect() as conn:
            conn.execute(
                "UPDATE users SET ppt_style = NULL, ppt_subject = NULL WHERE id = ?",
                (int(user["id"]),),
            )
            conn.commit()
        user = db.user_by_username("teacher")
        deck = Presentation()
        deck.slide_width = Inches(13.333)
        deck.slide_height = Inches(7.5)
        blank = deck.slide_layouts[6]
        cover = deck.slides.add_slide(blank)
        title = cover.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        run = title.text_frame.paragraphs[0].add_run()
        run.text = "旧课文不要出现"
        run.font.size = Pt(40)
        body = deck.slides.add_slide(blank)
        heading = body.shapes.add_textbox(Inches(1), Inches(0.6), Inches(8), Inches(1))
        heading_run = heading.text_frame.paragraphs[0].add_run()
        heading_run.text = "旧标题"
        heading_run.font.size = Pt(32)
        point = body.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
        point_run = point.text_frame.paragraphs[0].add_run()
        point_run.text = "旧要点"
        point_run.font.size = Pt(18)
        payload = BytesIO()
        deck.save(payload)
        store_pptx(int(user["id"]), "旧课件.pptx", payload.getvalue())
        self.assertIn("请选择", describe_style(user, conversation_id))
        self.assertIn("用我刚上传的那份课件样式", describe_style(user, conversation_id))
        uploaded, _subject = prepare_export(user, conversation_id, "upload", "")
        slides = [
            {"role": "课题与目标", "title": "认识分数", "bullets": ["能说出几分之一"], "source": "教案"},
            {"role": "导入", "title": "分月饼", "bullets": ["平均分成四份"], "source": "要求"},
            {"role": "新知", "title": "几分之一", "bullets": ["份数是分母"], "source": "课本"},
            {"role": "小结", "title": "今天记住", "bullets": ["分母是平均分的份数"], "source": "教案"},
            {"role": "作业", "title": "课后", "bullets": ["完成练习二"], "source": "课本"},
        ]
        saved = write_ppt(user, conversation_id, "认识分数", slides, uploaded, "数学")
        from app.config import DATA_DIR

        token = saved["url"].rsplit("/", 1)[-1]
        opened = Presentation(str(DATA_DIR / "docs" / f"{token}.pptx"))
        text = "\n".join(
            shape.text_frame.text
            for slide in opened.slides
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
        )
        self.assertIn("认识分数", text)
        self.assertIn("几分之一", text)
        self.assertNotIn("旧课文不要出现", text)
        self.assertNotIn("旧要点", text)
        with self.assertRaises(ValueError):
            store_pptx(int(user["id"]), "扫描.pdf", b"%PDF-1.4")

    def test_lesson_skills_leave_out_cloud_and_feishu(self):
        from app.agents.registry import skill_directories, tool_call_allowed
        names = [path.name for path in skill_directories("lesson")]
        self.assertEqual(names, ["exam-generator", "markdown-to-exam-docx"])
        self.assertEqual(skill_directories("office"), [])
        self.assertTrue(tool_call_allowed("export_ppt", "lesson"))
        self.assertFalse(tool_call_allowed("export_ppt", "office"))
        self.assertTrue(tool_call_allowed("ppt_style", "lesson"))
        self.assertFalse(tool_call_allowed("ppt_style", "office"))
        self.assertFalse(tool_call_allowed("export_exam", "office"))
        self.assertTrue(tool_call_allowed("export_materials", "office"))
        self.assertTrue(tool_call_allowed("save_mistakes", "lesson"))
        self.assertFalse(tool_call_allowed("save_mistakes", "office"))
        self.assertTrue(tool_call_allowed("review_parent_note", "office"))
        self.assertFalse(tool_call_allowed("review_parent_note", "lesson"))
        self.assertFalse(tool_call_allowed("bash"))
        self.assertFalse(tool_call_allowed(""))

    def test_router_sends_exam_and_notice_to_different_agents(self):
        from fastapi.testclient import TestClient

        from app.main import app

        seen = []

        async def fake_respond(user, identity, content=None, answers=None, skip=False, conversation_id=None):
            seen.append((identity, content))
            yield {"delta": identity}
            yield {"done": True, "message_id": 1}

        from unittest.mock import patch

        import app.agents.conductor as conductor
        from app.agents import router as agent_router
        from app.agents.router import decide

        agent_router.DEEPSEEK_API_KEY = ""
        self.assertEqual(decide("再出一道", "lesson")["route"], "lesson")
        mixed = decide("出一套试卷，再写一份家长通知", None)
        self.assertEqual(mixed["route"], "lesson")
        self.assertIn("通知", mixed["next"])

        client = TestClient(app)
        logged = client.post(
            "/login",
            data={"username": "teacher", "password": "test-password-1"},
            follow_redirects=False,
        )
        self.assertEqual(logged.status_code, 303)
        with patch.object(conductor, "_respond", fake_respond):
            exam_reply = client.post("/api/chat", json={"content": "帮我出一套期中试卷"})
            notice_reply = client.post("/api/chat", json={"content": "写一份家长会通知"})
        self.assertEqual(exam_reply.status_code, 200)
        self.assertEqual(notice_reply.status_code, 200)
        self.assertEqual(seen[0][0], "lesson")
        self.assertEqual(seen[1][0], "office")
        self.assertIn("试卷", seen[0][1])
        ask = client.post("/api/chat", json={"content": "你好"})
        self.assertEqual(ask.status_code, 200)
        self.assertIn("出题", ask.text)
        self.assertEqual(len(seen), 2)

    def test_office_cannot_export_an_exam(self):
        from app.exam import write_exam_files

        user = db.user_by_username("teacher")
        conversation_id = db.conversation_id(user["id"], "office")
        db.add_message(user["id"], "office", "user", "出一套题")
        with self.assertRaises(ValueError):
            write_exam_files(user, conversation_id, "学生卷", "【答案】B")

    def test_roster_import_keeps_an_existing_password(self):
        user = db.user_by_username("teacher")
        previous = user["password_hash"]
        token = db.create_session(user["id"], "lesson")
        try:
            db.import_accounts([
                {
                    "username": "teacher",
                    "display_name": "测试老师",
                    "password_hash": hash_password("roster-would-replace"),
                    "role": "admin",
                    "disabled": False,
                    "assigned_scopes": [],
                }
            ])
            updated = db.user_by_username("teacher")
            self.assertEqual(updated["password_hash"], previous)
            self.assertIsNotNone(db.session_user(token))
            db.import_accounts([
                {
                    "username": "teacher",
                    "display_name": "测试老师",
                    "password_hash": previous,
                    "role": "admin",
                    "disabled": True,
                    "assigned_scopes": [],
                }
            ])
            self.assertIsNone(db.session_user(token))
        finally:
            db.save_account(user["id"], "teacher", "测试老师", "admin", False, [], "")

    def test_expired_session_is_rejected(self):
        user = db.user_by_username("teacher")
        token = db.create_session(user["id"], "lesson")
        with db.connect() as conn:
            conn.execute(
                "UPDATE sessions SET expires_at = ? WHERE token = ?",
                ("2000-01-01T00:00:00+00:00", token),
            )
            conn.commit()
        self.assertIsNone(db.session_user(token))

    def test_disabling_an_account_drops_its_session(self):
        admin = db.user_by_username("teacher")
        db.create_account("lin2", "林老师", "password-ok", "teacher", [])
        lin = db.user_by_username("lin2")
        token = db.create_session(lin["id"], "lesson")
        db.save_account(admin["id"], "lin2", "林老师", "teacher", True, [], "")
        self.assertIsNone(db.session_user(token))


if __name__ == "__main__":
    unittest.main()
