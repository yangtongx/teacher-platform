#!/usr/bin/env python3
"""把结构化题目 JSON 导出为 xlsx 试卷文件。

退出码：0 成功 / 2 输入不合法 / 3 缺少依赖 / 4 写入失败。
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
except ImportError:  # pragma: no cover
    print(
        "错误：缺少 openpyxl。请安装：pip install openpyxl（或 apt-get install -y python3-openpyxl）",
        file=sys.stderr,
    )
    sys.exit(3)

COLUMNS = [
    ("题号", 6),
    ("题干", 60),
    ("答案", 24),
    ("解析", 60),
    ("题型", 12),
    ("难度", 8),
    ("知识点", 18),
]

REQUIRED = ("stem", "answer")
OPTIONAL = ("analysis", "type", "difficulty", "point")


def load_questions(path: Path) -> list:
    """读入并校验题目 JSON，任何不合法都抛 ValueError（消息面向老师可读）。"""
    if not path.is_file():
        raise ValueError(f"输入文件不存在：{path}")
    raw = path.read_text(encoding="utf-8")
    if not raw.strip():
        raise ValueError(f"输入文件是空的：{path}")
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ValueError(f"输入不是合法 JSON：{exc}") from exc
    if not isinstance(data, dict) or "questions" not in data:
        raise ValueError("JSON 顶层必须是对象，且含 questions 字段")
    questions = data["questions"]
    if not isinstance(questions, list) or not questions:
        raise ValueError("questions 必须是非空数组")
    for index, item in enumerate(questions, start=1):
        if not isinstance(item, dict):
            raise ValueError(f"第 {index} 题不是对象")
        for field in REQUIRED:
            value = item.get(field)
            if not isinstance(value, str) or not value.strip():
                raise ValueError(f"第 {index} 题缺少必填字段 {field}")
    return questions


def write_workbook(questions: list, out_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "试卷"

    for column, (title, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.font = Font(bold=True)
        sheet.column_dimensions[cell.column_letter].width = width

    for index, item in enumerate(questions, start=1):
        values = [index, item["stem"].strip(), item["answer"].strip()]
        for field in OPTIONAL:
            values.append(str(item.get(field) or "").strip())
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=index + 1, column=column, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    sheet.freeze_panes = "A2"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="把题目 JSON 导出为 xlsx")
    parser.add_argument("--input", required=True, help="题目 JSON 路径")
    parser.add_argument("--output", required=True, help="输出 xlsx 路径")
    args = parser.parse_args(argv)

    try:
        questions = load_questions(Path(args.input))
        out_path = Path(args.output)
        write_workbook(questions, out_path)
    except ValueError as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 2
    except OSError as exc:
        print(f"错误：写入失败：{exc}", file=sys.stderr)
        return 4

    print(out_path.resolve())
    return 0


if __name__ == "__main__":
    sys.exit(main())