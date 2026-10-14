#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""教案 Markdown → 教案 Word（docx）转换器。

契约来源：`workspaces/lesson/AGENTS.md` 阶段3

    python3 skills/lesson-plan-generator/scripts/convert_lesson.py \
      --input tmp/lesson.md --output lessons/{单元名}_教案.docx

Markdown 结构规范见 `skills/lesson-plan-generator/SKILL.md`（权威来源），本脚本
按同一规范解析：

    # 总标题（全文唯一）
    ## 单元教学设计            ← 单元目标 / 学情分析 / 单元评价量规 /
                                 实施计划（表格）/ 学习活动 / 单元作业
    ## 第N课时 — {课题名称}（{课型}）
       教学分析 / 学习目标 / 学习过程（表格或分步小标题）/ 设计意图 /
       板书设计（不能为空）/ 作业与拓展

只用 python-docx + 标准库实现，不引入任何新依赖。中文字体同时设置
`run.font.name` 与 `rPr/rFonts/@w:eastAsia`，否则中文会退化成默认字体。
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError:  # pragma: no cover - 仅在服务器缺依赖时触发
    sys.stderr.write(
        "错误：缺少依赖 python-docx。请先安装（pip install python-docx），再运行本脚本。\n"
    )
    raise SystemExit(2)

# ---------------------------------------------------------------- 排版常量

FONT_CJK = "宋体"          # 中文默认字体（契约：正文宋体）
FONT_LATIN = "Times New Roman"  # 西文默认字体
SIZE_BODY = 10.5           # 五号字 = 10.5pt
SIZE_TITLE = 18.0          # 教案总标题
SIZE_LESSON = 14.0         # 课时标题
SIZE_SECTION = 12.0        # 单元小节 / 课时内字段
SIZE_STEP = 11.0           # 学习过程里的分步小标题
LINE_SPACING = 1.5         # 契约：1.5 倍行距
FIRST_LINE_INDENT = Cm(0.74)  # 正文首行缩进约 2 字符

# 页面：A4 + 上下 2.54cm、左右 3.17cm
PAGE_W, PAGE_H = Cm(21.0), Cm(29.7)
MARGIN_TB, MARGIN_LR = Cm(2.54), Cm(3.17)

# 单元教学设计的必备子字段（AGENTS.md 阶段2 第1步）
UNIT_FIELDS = ["单元目标", "学情分析", "单元评价量规", "实施计划", "学习活动", "单元作业"]
# 课时必备字段（AGENTS.md 阶段2 第2步）
LESSON_FIELDS = ["教学分析", "学习目标", "学习过程", "设计意图", "板书设计", "作业与拓展"]

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
TABLE_SEP_CELL_RE = re.compile(r"^:?-+:?$")
ORDERED_RE = re.compile(r"^\s*(\d+)[.、)]\s+(.*)$")
UNORDERED_RE = re.compile(r"^\s*[-*+]\s+(.*)$")
BOLD_ONLY_RE = re.compile(r"^\*\*(.+?)\*\*[：:]\s*$")
STARS_RE = re.compile(r"\*\*(.+?)\*\*")


def configure_std_streams() -> None:
    """Windows 控制台默认 GBK，中文与特殊符号会报 UnicodeEncodeError。"""
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
        except Exception:
            pass


# ---------------------------------------------------------------- 字体工具


def set_run_font(run, cjk: str = FONT_CJK, latin: str = FONT_LATIN) -> None:
    """同时设置西文字体与 `w:eastAsia`，否则中文不生效。"""
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)
    rfonts.set(qn("w:cs"), latin)


def add_runs(paragraph, text: str, *, cjk: str, latin: str, size: float,
             bold: bool = False) -> None:
    """按 `**加粗**` 把一行文本拆成多个 run，逐段施加字体与字号。"""
    pos = 0
    for match in STARS_RE.finditer(text):
        if match.start() > pos:
            _plain_run(paragraph, text[pos:match.start()], cjk, latin, size, bold)
        _plain_run(paragraph, match.group(1), cjk, latin, size, True)
        pos = match.end()
    if pos < len(text):
        _plain_run(paragraph, text[pos:], cjk, latin, size, bold)


def _plain_run(paragraph, text: str, cjk: str, latin: str, size: float,
               bold: bool) -> None:
    if text == "":
        return
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    set_run_font(run, cjk, latin)


def style_paragraph(paragraph, *, size: float, line_spacing: float = LINE_SPACING,
                    align=None, indent=None, space_after: float = 3.0,
                    space_before: float = 0.0) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing = line_spacing
    fmt.space_after = Pt(space_after)
    fmt.space_before = Pt(space_before)
    if align is not None:
        paragraph.alignment = align
    if indent is not None:
        fmt.first_line_indent = indent


def setup_document(*, cjk: str, latin: str, body_size: float,
                   line_spacing: float) -> "Document":
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = latin
    style.font.size = Pt(body_size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)
    rfonts.set(qn("w:cs"), latin)
    for section in doc.sections:
        section.page_width = PAGE_W
        section.page_height = PAGE_H
        section.top_margin = MARGIN_TB
        section.bottom_margin = MARGIN_TB
        section.left_margin = MARGIN_LR
        section.right_margin = MARGIN_LR
    return doc


# ---------------------------------------------------------------- 解析

def split_table_row(line: str) -> list:
    """把 `| a | b |` 拆成单元格文本；`\\|` 表示字面竖线。"""
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    cells = []
    buf = ""
    i = 0
    while i < len(text):
        ch = text[i]
        if ch == "\\" and i + 1 < len(text) and text[i + 1] == "|":
            buf += "|"
            i += 2
            continue
        if ch == "|":
            cells.append(buf.strip())
            buf = ""
            i += 1
            continue
        buf += ch
        i += 1
    cells.append(buf.strip())
    return cells


def is_table_separator(cells: list) -> bool:
    if not cells:
        return False
    return all(TABLE_SEP_CELL_RE.match(cell.replace(" ", "")) for cell in cells)


def parse_blocks(text: str) -> list:
    """按行解析为块列表；表格与段落/列表不会互相吞并。"""
    blocks = []
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    i = 0
    while i < len(lines):
        raw = lines[i]
        # 去掉 UTF-8 BOM 与零宽字符，否则首行 `# 标题` 认不出来
        line = raw.strip().lstrip("\ufeff\u200b")
        if line == "":
            i += 1
            continue
        if line.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1  # 跳过结束围栏
            blocks.append({"type": "code", "text": "\n".join(code)})
            continue
        heading = HEADING_RE.match(line)
        if heading:
            blocks.append({
                "type": "heading",
                "level": len(heading.group(1)),
                "text": heading.group(2).strip(),
            })
            i += 1
            continue
        if "|" in line and i + 1 < len(lines):
            head_cells = split_table_row(line)
            sep_cells = split_table_row(lines[i + 1])
            if len(head_cells) >= 2 and len(sep_cells) == len(head_cells) \
                    and is_table_separator(sep_cells):
                rows = [head_cells]
                i += 2
                while i < len(lines) and "|" in lines[i] and lines[i].strip():
                    cells = split_table_row(lines[i])
                    if len(cells) < len(head_cells):
                        cells = cells + [""] * (len(head_cells) - len(cells))
                    rows.append(cells[:len(head_cells)])
                    i += 1
                blocks.append({"type": "table", "rows": rows})
                continue
        ordered = ORDERED_RE.match(raw)
        if ordered:
            blocks.append({"type": "ol", "number": ordered.group(1),
                           "text": ordered.group(2).strip()})
            i += 1
            continue
        unordered = UNORDERED_RE.match(raw)
        if unordered:
            blocks.append({"type": "ul", "text": unordered.group(1).strip()})
            i += 1
            continue
        bold_only = BOLD_ONLY_RE.match(line)
        if bold_only:
            blocks.append({"type": "boldline", "text": bold_only.group(1).strip()})
            i += 1
            continue
        blocks.append({"type": "para", "text": line})
        i += 1
    return blocks


def same_field(text: str, field: str) -> bool:
    cleaned = text.replace("：", "").replace(":", "").replace(" ", "").strip()
    cleaned = cleaned.replace("*", "")
    return cleaned == field


def group_unit(blocks: list) -> list:
    """把 `## 单元教学设计` 下面的块按字段分组，缺失字段补 None。"""
    groups = []
    current = None
    for block in blocks:
        key = None
        if block["type"] == "heading" and block["level"] == 3:
            key = block["text"]
        elif block["type"] == "boldline":
            key = block["text"]
        if key is not None:
            matched = next((f for f in UNIT_FIELDS if same_field(key, f)), None)
            if matched is not None:
                current = {"field": matched, "blocks": []}
                groups.append(current)
                continue
        if current is None:
            current = {"field": None, "blocks": []}
            groups.append(current)
        current["blocks"].append(block)
    present = {g["field"] for g in groups if g["field"]}
    for field in UNIT_FIELDS:
        if field not in present:
            groups.append({"field": field, "blocks": [], "missing": True})
    return groups


def group_lessons(blocks: list) -> list:
    """切分课时：每个 `## ` 标题起一个新课时，内部按字段分组。"""
    lessons = []
    current = None
    for block in blocks:
        if block["type"] == "heading" and block["level"] == 2:
            current = {"title": block["text"], "groups": [], "fields": {}}
            lessons.append(current)
            continue
        if current is None:
            current = {"title": None, "groups": [], "fields": {}}
            lessons.append(current)
        key = None
        if block["type"] == "heading" and block["level"] == 3:
            key = block["text"]
        elif block["type"] == "boldline":
            key = block["text"]
        if key is not None:
            matched = next((f for f in LESSON_FIELDS if same_field(key, f)), None)
            if matched is not None:
                group = {"field": matched, "blocks": []}
                current["groups"].append(group)
                current["fields"][matched] = group
                continue
        if not current["groups"] or current["groups"][-1]["field"] is None:
            current["groups"].append({"field": None, "blocks": [block]})
        else:
            current["groups"][-1]["blocks"].append(block)
    return lessons


# ---------------------------------------------------------------- 渲染

def render_heading(doc, text, *, level, cjk, latin, body_size, line_spacing):
    if level == 1:
        para = doc.add_paragraph()
        style_paragraph(para, size=SIZE_TITLE, line_spacing=line_spacing,
                        align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10.0,
                        space_before=0.0)
        add_runs(para, text, cjk=cjk, latin=latin, size=SIZE_TITLE, bold=True)
    elif level == 2:
        para = doc.add_paragraph()
        style_paragraph(para, size=SIZE_LESSON, line_spacing=line_spacing,
                        align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6.0,
                        space_before=10.0)
        add_runs(para, text, cjk=cjk, latin=latin, size=SIZE_LESSON, bold=True)
    elif level == 3:
        para = doc.add_paragraph()
        style_paragraph(para, size=SIZE_SECTION, line_spacing=line_spacing,
                        space_after=4.0, space_before=6.0)
        add_runs(para, text, cjk=cjk, latin=latin, size=SIZE_SECTION, bold=True)
    else:
        para = doc.add_paragraph()
        style_paragraph(para, size=SIZE_STEP, line_spacing=line_spacing,
                        space_after=3.0, space_before=4.0)
        add_runs(para, text, cjk=cjk, latin=latin, size=SIZE_STEP, bold=True)
    return para


def render_table(doc, rows, *, cjk, latin, body_size, line_spacing):
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_index, row in enumerate(rows):
        for c_index in range(cols):
            cell = table.cell(r_index, c_index)
            text = row[c_index] if c_index < len(row) else ""
            para = cell.paragraphs[0]
            style_paragraph(para, size=body_size, line_spacing=1.0,
                            align=WD_ALIGN_PARAGRAPH.CENTER if r_index == 0
                            else WD_ALIGN_PARAGRAPH.LEFT,
                            space_after=0.0)
            add_runs(para, text, cjk=cjk, latin=latin, size=body_size,
                     bold=(r_index == 0))
    spacer = doc.add_paragraph()
    style_paragraph(spacer, size=body_size, line_spacing=1.0, space_after=0.0)
    return table


def render_blocks(doc, blocks, *, cjk, latin, body_size, line_spacing,
                  first_level=4) -> None:
    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            level = min(block["level"] + first_level - 1, 6)
            render_heading(doc, block["text"], level=level, cjk=cjk, latin=latin,
                           body_size=body_size, line_spacing=line_spacing)
        elif kind == "boldline":
            para = doc.add_paragraph()
            style_paragraph(para, size=SIZE_STEP, line_spacing=line_spacing,
                            space_after=3.0, space_before=4.0)
            add_runs(para, block["text"], cjk=cjk, latin=latin,
                     size=SIZE_STEP, bold=True)
        elif kind == "table":
            render_table(doc, block["rows"], cjk=cjk, latin=latin,
                         body_size=body_size, line_spacing=line_spacing)
        elif kind == "para":
            para = doc.add_paragraph()
            style_paragraph(para, size=body_size, line_spacing=line_spacing,
                            indent=FIRST_LINE_INDENT)
            add_runs(para, block["text"], cjk=cjk, latin=latin, size=body_size)
        elif kind == "ol":
            para = doc.add_paragraph()
            style_paragraph(para, size=body_size, line_spacing=line_spacing,
                            space_after=2.0)
            fmt = para.paragraph_format
            fmt.left_indent = Cm(0.74)
            fmt.first_line_indent = Cm(-0.74)
            add_runs(para, f"{block['number']}. {block['text']}", cjk=cjk,
                     latin=latin, size=body_size)
        elif kind == "ul":
            para = doc.add_paragraph()
            style_paragraph(para, size=body_size, line_spacing=line_spacing,
                            space_after=2.0)
            fmt = para.paragraph_format
            fmt.left_indent = Cm(0.74)
            fmt.first_line_indent = Cm(-0.74)
            add_runs(para, f"• {block['text']}", cjk=cjk, latin=latin,
                     size=body_size)
        elif kind == "code":
            for line in block["text"].split("\n"):
                para = doc.add_paragraph()
                style_paragraph(para, size=body_size, line_spacing=1.0,
                                space_after=0.0)
                add_runs(para, line, cjk=cjk, latin=latin, size=body_size)


def render_missing_field(doc, field, *, cjk, latin, body_size, line_spacing):
    para = doc.add_paragraph()
    style_paragraph(para, size=SIZE_SECTION, line_spacing=line_spacing,
                    space_after=4.0, space_before=6.0)
    add_runs(para, field, cjk=cjk, latin=latin, size=SIZE_SECTION, bold=True)
    placeholder = doc.add_paragraph()
    style_paragraph(placeholder, size=body_size, line_spacing=line_spacing,
                    indent=FIRST_LINE_INDENT)
    add_runs(placeholder, "（本字段缺失，请老师补充）", cjk=cjk, latin=latin,
             size=body_size)


# ---------------------------------------------------------------- 校验

def check_lesson(text: str) -> list:
    """返回警告列表；板书设计缺失（或为空）必须给出警告。"""
    warnings = []
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    title_lines = []
    for raw in text.split("\n"):
        match = HEADING_RE.match(raw.strip().lstrip("\ufeff\u200b"))
        if match and len(match.group(1)) == 1:
            title_lines.append(match.group(2).strip())
    if not title_lines:
        warnings.append("缺少 `#` 总标题：教案必须有且只有一个一级标题")
    elif len(title_lines) > 1:
        warnings.append(
            f"检测到 {len(title_lines)} 个 `#` 总标题，教案总标题必须全文唯一"
        )

    blocks = parse_blocks(text)
    has_unit = any(b["type"] == "heading" and b["level"] == 2
                   and same_field(b["text"], "单元教学设计") for b in blocks)
    if not has_unit:
        warnings.append("缺少 `## 单元教学设计` 小节")
    else:
        # 单元子字段既允许 `### 单元目标`，也允许 `**单元目标：**`，两种都认
        for field in UNIT_FIELDS:
            if not field_present(blocks, field):
                warnings.append(f"「单元教学设计」缺少「{field}」字段")

    lesson_blocks = group_lessons(blocks)
    real_lessons = [l for l in lesson_blocks if l["title"]]
    if not real_lessons:
        warnings.append("没有检测到任何 `## 第N课时 — {课题名称}（{课型}）` 课时标题")
    for lesson in real_lessons:
        title = lesson["title"]
        if not re.match(r"^第\s*\d+\s*课时", title):
            # 「单元教学设计」小节不是课时，不能按课时字段去要求它
            if same_field(title, "单元教学设计"):
                continue
            warnings.append(
                f"课时标题「{title}」不符合 `## 第N课时 — {{课题名称}}（{{课型}}）` 规范"
            )
            continue
        board = lesson["fields"].get("板书设计")
        if board is None:
            warnings.append(f"课时「{title}」缺少「板书设计」字段")
        elif not any(b["type"] != "heading" for b in board["blocks"]):
            warnings.append(f"课时「{title}」的「板书设计」为空，必须补全")
        for field in LESSON_FIELDS:
            if field not in lesson["fields"] and field != "板书设计":
                # 「板书设计」缺失已在上面单独告警，避免重复
                warnings.append(f"课时「{title}」缺少「{field}」字段")
    return warnings


def field_present(blocks: list, field: str) -> bool:
    for block in blocks:
        if block["type"] == "heading" and block["level"] == 3 \
                and same_field(block["text"], field):
            return True
        if block["type"] == "boldline" and same_field(block["text"], field):
            return True
    return False


# ---------------------------------------------------------------- 主流程

def count_words(text: str) -> int:
    """字数统计：CJK 按字计数 + 西文按词计数。"""
    cjk = len(re.findall(r"[\u3400-\u4dbf\u4e00-\u9fff\u3040-\u30ff\uac00-\ud7af]", text))
    words = len(re.findall(r"[A-Za-z0-9]+", text))
    return cjk + words


def build_document(text: str, *, cjk: str, latin: str, body_size: float,
                   line_spacing: float) -> "Document":
    doc = setup_document(cjk=cjk, latin=latin, body_size=body_size,
                         line_spacing=line_spacing)
    blocks = parse_blocks(text)

    # 1) `#` 总标题：只取第一个
    title_done = False
    rest = []
    for block in blocks:
        if not title_done and block["type"] == "heading" and block["level"] == 1:
            render_heading(doc, block["text"], level=1, cjk=cjk, latin=latin,
                           body_size=body_size, line_spacing=line_spacing)
            title_done = True
            continue
        rest.append(block)

    # 2) 切出 `## 单元教学设计` 与各课时：
    #    `##` 级标题即分界线，它下面的正文归属于它——单元小节下面的内容
    #    属于单元部分，不能一律塞进课时段。
    unit_blocks, lesson_section = [], []
    in_unit = False      # 当前是否位于「单元教学设计」小节内
    seen_unit = False    # 是否已经遇到「单元教学设计」
    for block in rest:
        is_h2 = block["type"] == "heading" and block["level"] == 2
        if is_h2:
            if same_field(block["text"], "单元教学设计") and not seen_unit:
                seen_unit = True
                in_unit = True
                unit_blocks.append(block)
                continue
            in_unit = False
            lesson_section.append(block)
            continue
        if in_unit:
            unit_blocks.append(block)
        else:
            lesson_section.append(block)

    # 3) `## 单元教学设计` 子字段：缺失的补占位，保证结构完整
    unit_body = []
    for block in unit_blocks:
        if block["type"] == "heading" and block["level"] == 2:
            render_heading(doc, block["text"], level=2, cjk=cjk, latin=latin,
                           body_size=body_size, line_spacing=line_spacing)
            continue
        unit_body.append(block)
    for group in group_unit(unit_body):
        if group["field"] is None:
            render_blocks(doc, group["blocks"], cjk=cjk, latin=latin,
                          body_size=body_size, line_spacing=line_spacing,
                          first_level=4)
            continue
        render_heading(doc, group["field"], level=3, cjk=cjk, latin=latin,
                       body_size=body_size, line_spacing=line_spacing)
        if group.get("missing"):
            placeholder = doc.add_paragraph()
            style_paragraph(placeholder, size=body_size,
                            line_spacing=line_spacing,
                            indent=FIRST_LINE_INDENT)
            add_runs(placeholder, "（本字段缺失，请老师补充）", cjk=cjk,
                     latin=latin, size=body_size)
            continue
        render_blocks(doc, group["blocks"], cjk=cjk, latin=latin,
                      body_size=body_size, line_spacing=line_spacing,
                      first_level=4)

    # 4) 各课时
    for lesson in group_lessons(lesson_section):
        if lesson["title"]:
            render_heading(doc, lesson["title"], level=2, cjk=cjk, latin=latin,
                           body_size=body_size, line_spacing=line_spacing)
        for field in LESSON_FIELDS:
            group = lesson["fields"].get(field)
            if group is None:
                render_missing_field(doc, field, cjk=cjk, latin=latin,
                                     body_size=body_size,
                                     line_spacing=line_spacing)
                continue
            render_heading(doc, field, level=3, cjk=cjk, latin=latin,
                           body_size=body_size, line_spacing=line_spacing)
            render_blocks(doc, group["blocks"], cjk=cjk, latin=latin,
                          body_size=body_size, line_spacing=line_spacing,
                          first_level=4)
        for group in lesson["groups"]:
            if group["field"] is None and group["blocks"]:
                render_blocks(doc, group["blocks"], cjk=cjk, latin=latin,
                              body_size=body_size, line_spacing=line_spacing,
                              first_level=4)
    return doc


def main(argv=None) -> int:
    configure_std_streams()
    parser = argparse.ArgumentParser(
        prog="convert_lesson.py",
        description="把教案 Markdown（tmp/lesson.md）转换为教案 Word 文档。",
        epilog="示例：python3 skills/lesson-plan-generator/scripts/convert_lesson.py "
               "--input tmp/lesson.md --output lessons/第一单元_教案.docx",
    )
    parser.add_argument("--input", required=True, help="教案 Markdown 路径，如 tmp/lesson.md")
    parser.add_argument("--output", required=True, help="输出 docx 路径，如 lessons/{单元名}_教案.docx")
    parser.add_argument("--font-cjk", default=FONT_CJK,
                        help=f"中文字体（默认 {FONT_CJK}）")
    parser.add_argument("--font-latin", default=FONT_LATIN,
                        help=f"西文字体（默认 {FONT_LATIN}）")
    parser.add_argument("--body-size", type=float, default=SIZE_BODY,
                        help=f"正文字号，单位 pt（默认 {SIZE_BODY}，即五号）")
    parser.add_argument("--line-spacing", type=float, default=LINE_SPACING,
                        help=f"行距倍数（默认 {LINE_SPACING}）")
    parser.add_argument("--no-warn", action="store_true",
                        help="不打印格式警告（板书设计缺失等）")
    args = parser.parse_args(argv)

    if args.body_size <= 0:
        sys.stderr.write(f"错误：--body-size 必须为正数，收到 {args.body_size}\n")
        return 2
    if args.line_spacing <= 0:
        sys.stderr.write(f"错误：--line-spacing 必须为正数，收到 {args.line_spacing}\n")
        return 2

    src = Path(args.input)
    if not src.is_file():
        sys.stderr.write(f"错误：输入文件不存在：{src}\n")
        return 2
    try:
        text = src.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        sys.stderr.write(f"错误：输入文件不是 UTF-8 编码，无法解析：{src}（{exc}）\n")
        return 2
    if not text.strip():
        sys.stderr.write(f"错误：输入文件为空：{src}。教案 Markdown 不能是空文件。\n")
        return 2

    warnings = check_lesson(text)
    if warnings and not args.no_warn:
        for item in warnings:
            sys.stderr.write(f"警告：{item}\n")

    doc = build_document(text, cjk=args.font_cjk, latin=args.font_latin,
                         body_size=args.body_size,
                         line_spacing=args.line_spacing)
    out = Path(args.output)
    try:
        if out.parent and str(out.parent) not in ("", "."):
            out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))
    except OSError as exc:
        sys.stderr.write(f"错误：写入 docx 失败：{out}（{exc}）\n")
        return 1

    words = count_words(text)
    print(str(out.resolve()))
    print(f"已生成教案：{out.resolve()}（相对路径：{out}）")
    print(f"字数统计：{words} 字（中文字符 + 西文单词）；警告 {len(warnings)} 条")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
