import os
import sqlite3
import tempfile
import unittest
from pathlib import Path

os.environ["DATA_DIR"] = tempfile.mkdtemp()
os.environ["ADMIN_USERNAME"] = "teacher"
os.environ["ADMIN_PASSWORD"] = "test-password-1"
os.environ["ADMIN_DISPLAY_NAME"] = "测试老师"

from app.auth import hash_password, verify_password
from app import db
from app.corpus import textbook_dbs
from app.rag import search


class AuthTests(unittest.TestCase):
    def test_roundtrip_matches_existing_format(self):
        stored = hash_password("correct-horse")
        self.assertTrue(stored.startswith("pbkdf2_sha256$600000$"))
        self.assertTrue(verify_password("correct-horse", stored))
        self.assertFalse(verify_password("wrong-horse", stored))
        self.assertFalse(verify_password("correct-horse", "not-a-hash"))

    def test_switch_keeps_two_threads(self):
        db.init_db()
        user = db.user_by_username("teacher")
        token = db.create_session(user["id"], "lesson")
        db.add_message(user["id"], "lesson", "user", "出一套题")
        db.set_identity(token, "office")
        _, identity = db.session_user(token)
        self.assertEqual(identity, "office")
        self.assertEqual(db.list_messages(user["id"], "office"), [])
        lesson = db.list_messages(user["id"], "lesson")
        self.assertEqual(lesson[0]["content"], "出一套题")
        db.set_identity(token, "lesson")
        _, identity = db.session_user(token)
        self.assertEqual(identity, "lesson")

    def test_import_roster_hash_and_scope(self):
        db.init_db()
        stored = hash_password("roster-secret")
        imported = db.import_accounts([
            {
                "username": "zhang",
                "display_name": "张老师",
                "password_hash": stored,
                "role": "teacher",
                "disabled": False,
                "stage": "junior",
                "subject": "chinese",
                "textbook_version": "tb",
                "grade": "8",
                "assigned_scopes": ["junior/chinese/tb/8上"],
            }
        ])
        self.assertEqual(imported, 1)
        user = db.user_by_username("zhang")
        self.assertTrue(verify_password("roster-secret", user["password_hash"]))
        self.assertEqual(user["subject"], "chinese")
        root = Path(tempfile.mkdtemp())
        book = root / "junior" / "chinese" / "tb" / "8上" / "rag"
        book.mkdir(parents=True)
        (book / "textbook.db").write_bytes(b"")
        found = textbook_dbs(user, root)
        self.assertEqual(len(found), 1)
        self.assertEqual(found[0]["id"], "junior/chinese/tb/8上")

    def test_search_uses_same_bigram_index(self):
        root = Path(tempfile.mkdtemp())
        db_path = root / "textbook.db"
        content = "分数的意义是把单位一平均分成若干份"
        from app.rag import tokenize

        tokens = [tok for tok in tokenize(content) if tok != "\n"]
        conn = sqlite3.connect(db_path)
        conn.execute(
            """
            CREATE TABLE chunks (
              chunk_id INTEGER PRIMARY KEY,
              source_file TEXT,
              heading TEXT,
              content TEXT,
              chunk_index INTEGER,
              page_start INTEGER,
              page_end INTEGER
            )
            """
        )
        conn.execute(
            """
            CREATE VIRTUAL TABLE fts_chunks USING fts5(
              content, heading, content='chunks', content_rowid='chunk_id', tokenize='unicode61'
            )
            """
        )
        conn.execute(
            "INSERT INTO chunks VALUES (1, '数学.pdf', '分数', ?, 0, 12, 13)",
            (content,),
        )
        conn.execute(
            "INSERT INTO fts_chunks(rowid, content, heading) VALUES (1, ?, ?)",
            (" ".join(tokens), "分数"),
        )
        conn.commit()
        conn.close()
        hits = search(db_path, "分数的意义")
        self.assertEqual(len(hits), 1)
        self.assertIn("平均分", hits[0]["content"])
        self.assertEqual(hits[0]["page_start"], 12)

    def test_word_file_belongs_to_the_author(self):
        from docx import Document

        from app.docx import document_name, write_docx

        db.init_db()
        owner = db.user_by_username("teacher")
        db.add_message(owner["id"], "office", "assistant", "# 放假通知\n\n明天下午放假。")
        message = db.list_messages(owner["id"], "office")[-1]
        self.assertEqual(document_name(message["content"]), "放假通知")
        path = Path(os.environ["DATA_DIR"]) / "docs" / "tok.docx"
        write_docx(message["content"], path)
        text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
        self.assertIn("放假通知", text)
        self.assertIn("明天下午放假", text)
        db.save_document(owner["id"], message["id"], "tok", "放假通知.docx")
        self.assertEqual(db.document_for_user(owner["id"], "tok")["filename"], "放假通知.docx")
        self.assertIsNone(db.document_for_user(owner["id"] + 99, "tok"))
        self.assertIsNone(db.message_for_user(owner["id"] + 99, message["id"]))

    def test_password_can_be_replaced(self):
        db.init_db()
        user = db.user_by_username("teacher")
        db.set_password_hash(user["id"], hash_password("new-password-2"))
        updated = db.user_by_username("teacher")
        self.assertTrue(verify_password("new-password-2", updated["password_hash"]))
        self.assertFalse(verify_password("test-password-1", updated["password_hash"]))


class ConversationTests(unittest.TestCase):
    def test_new_conversation_keeps_the_previous_one(self):
        db.init_db()
        user = db.user_by_username("teacher")
        db.add_message(user["id"], "office", "user", "写一份放假通知")
        previous = db.conversation_id(user["id"], "office")
        self.assertEqual(db.list_conversations(user["id"], "office")[0]["title"], "写一份放假通知")
        opened = db.create_conversation(user["id"], "office")
        self.assertNotEqual(opened, previous)
        self.assertEqual(db.list_messages(user["id"], "office"), [])
        self.assertTrue(db.open_conversation(user["id"], "office", previous))
        restored = [item["content"] for item in db.list_messages(user["id"], "office")]
        self.assertIn("写一份放假通知", restored)
        self.assertFalse(db.open_conversation(user["id"], "lesson", previous))

    def test_upload_stays_on_its_conversation(self):
        from app.extract import extract_text

        db.init_db()
        user = db.user_by_username("teacher")
        text = extract_text("note.txt", "周五下午开家长会".encode("utf-8"))
        db.add_upload(user["id"], "lesson", "note.txt", text)
        self.assertIn("家长会", db.upload_context(user["id"], "lesson"))
        self.assertEqual(db.upload_context(user["id"], "office"), "")
        db.create_conversation(user["id"], "lesson")
        self.assertEqual(db.upload_context(user["id"], "lesson"), "")

    def test_docx_text_is_readable(self):
        from docx import Document

        from app.extract import extract_text

        path = Path(tempfile.mkdtemp()) / "lesson.docx"
        document = Document()
        document.add_paragraph("第三单元练习")
        document.save(path)
        self.assertIn("第三单元", extract_text(path.name, path.read_bytes()))
        with self.assertRaises(ValueError):
            extract_text("sheet.xlsx", b"not-a-doc")

    def test_old_unique_thread_can_split(self):
        folder = Path(tempfile.mkdtemp())
        previous_dir, previous_path = db.DATA_DIR, db.DB_PATH
        db.DATA_DIR = folder
        db.DB_PATH = folder / "app.db"
        try:
            conn = sqlite3.connect(db.DB_PATH)
            conn.executescript(
                """
                CREATE TABLE users (
                  id INTEGER PRIMARY KEY,
                  username TEXT NOT NULL UNIQUE,
                  display_name TEXT NOT NULL,
                  password_hash TEXT NOT NULL,
                  role TEXT NOT NULL DEFAULT 'teacher',
                  disabled INTEGER NOT NULL DEFAULT 0,
                  created_at TEXT NOT NULL
                );
                CREATE TABLE conversations (
                  id INTEGER PRIMARY KEY,
                  user_id INTEGER NOT NULL,
                  identity TEXT NOT NULL,
                  created_at TEXT NOT NULL,
                  UNIQUE(user_id, identity)
                );
                CREATE TABLE messages (
                  id INTEGER PRIMARY KEY,
                  conversation_id INTEGER NOT NULL,
                  role TEXT NOT NULL,
                  content TEXT NOT NULL,
                  created_at TEXT NOT NULL
                );
                INSERT INTO users (username, display_name, password_hash, role, created_at)
                  VALUES ('teacher', '测试老师', 'hash', 'teacher', '2026-01-01T00:00:00+00:00');
                INSERT INTO conversations (user_id, identity, created_at)
                  VALUES (1, 'office', '2026-01-01T00:00:00+00:00');
                INSERT INTO messages (conversation_id, role, content, created_at)
                  VALUES (1, 'user', '整理会议纪要', '2026-01-01T00:00:00+00:00');
                """
            )
            conn.commit()
            conn.close()
            db.init_db()
            user = db.user_by_username("teacher")
            self.assertEqual(user["role"], "admin")
            self.assertEqual(db.list_messages(user["id"], "office")[0]["content"], "整理会议纪要")
            self.assertEqual(db.list_conversations(user["id"], "office")[0]["title"], "整理会议纪要")
            db.create_conversation(user["id"], "office")
            self.assertEqual(len(db.list_conversations(user["id"], "office")), 2)
        finally:
            db.DATA_DIR = previous_dir
            db.DB_PATH = previous_path

    def test_admin_can_assign_books_and_disable(self):
        from app import corpus
        from app.corpus import account_scope_ids

        db.init_db()
        root = Path(tempfile.mkdtemp())
        scope = "junior/chinese/tb/8"
        book = root / "junior" / "chinese" / "tb" / "8" / "rag"
        book.mkdir(parents=True)
        (book / "textbook.db").write_bytes(b"")
        (root / "curriculum-catalog.json").write_text(
            '{"scopes":[{"id":"junior/chinese/tb/8","label":"8"}]}',
            encoding="utf-8",
        )
        previous = corpus.CORPUS_ROOT
        corpus.CORPUS_ROOT = root
        try:
            admin = db.user_by_username("teacher")
            db.set_password_hash(admin["id"], hash_password("test-password-1"))
            db.create_account("lin", "林老师", "password-ok", "teacher", [])
            with self.assertRaises(ValueError):
                db.save_account(admin["id"], "teacher", "测试老师", "admin", True, [], "")
            db.save_account(admin["id"], "lin", "林老师", "teacher", False, [scope], "")
            lin = db.user_by_username("lin")
            self.assertEqual(textbook_dbs(lin, root)[0]["id"], scope)
            db.save_account(admin["id"], "lin", "林老师", "teacher", True, [], "")
            lin = db.user_by_username("lin")
            self.assertEqual(lin["disabled"], 1)
            self.assertEqual(account_scope_ids(lin), [])

            from fastapi.testclient import TestClient

            from app.main import app

            client = TestClient(app)
            denied = client.post("/api/accounts", json={"username": "lin"})
            self.assertEqual(denied.status_code, 401)
            logged = client.post(
                "/login",
                data={"username": "teacher", "password": "test-password-1"},
                follow_redirects=False,
            )
            self.assertEqual(logged.status_code, 303)
            page = client.get("/api/accounts")
            self.assertEqual(page.status_code, 200)
            names = [item["username"] for item in page.json()["accounts"]]
            self.assertIn("lin", names)
            uploaded = client.post(
                "/api/uploads",
                files={"file": ("note.txt", "请把这篇通知读一遍".encode("utf-8"), "text/plain")},
            )
            self.assertEqual(uploaded.status_code, 200)
            self.assertEqual(uploaded.json()["uploads"][0]["filename"], "note.txt")
            rejected = client.post(
                "/api/uploads",
                files={"file": ("note.exe", b"abc", "application/octet-stream")},
            )
            self.assertEqual(rejected.status_code, 400)
        finally:
            corpus.CORPUS_ROOT = previous


class ExamFlowTests(unittest.TestCase):
    def test_senior_english_asks_listening_first(self):
        from app.exam import questions_for

        books = [{"id": "senior/english/rj/必修3", "label": "必修3", "path": Path(".")}]
        questions = questions_for(True, books)
        self.assertEqual(questions[0]["id"], "listening")
        self.assertIn("120分", questions[0]["options"][0]["label"])
        self.assertEqual(questions[1]["id"], "difficulty")
        self.assertIn("3:5:2", questions[1]["options"][0]["id"])

    def test_workspace_skill_and_two_docx_export(self):
        from io import BytesIO

        from docx import Document
        from fastapi.testclient import TestClient

        from app.config import ROOT
        from app.exam import student_markdown, write_exam_files
        from app.main import app

        skill = ROOT / "workspaces" / "lesson" / "skills" / "exam-generator" / "SKILL.md"
        self.assertTrue(skill.is_file(), "出卷技能应随工作区放在本项目里")
        self.assertNotIn("【答案】", student_markdown("1.（3分）题干\n【答案】B"))
        db.init_db()
        user = db.user_by_username("teacher")
        conversation_id = db.conversation_id(user["id"], "lesson")
        db.add_message(user["id"], "lesson", "user", "出一套题")
        student = (
            "# 八年级数学期中试卷（满分：3分 时间：20分钟）\n\n"
            "姓名：__________　班级：__________　得分：__________\n\n"
            "## 一、单项选择题（每题3分，共3分）\n\n"
            "1.（3分）下列各数中，是无理数的是（  ）\n"
            "A. 0.5\nB. √2\nC. 1/3\nD. -7\n"
        )
        answer = (
            "# 八年级数学期中试卷（满分：3分 时间：20分钟）\n\n"
            "## 一、单项选择题（每题3分，共3分）\n\n"
            "1.（3分）下列各数中，是无理数的是（  ）\n"
            "A. 0.5\nB. √2\nC. 1/3\nD. -7\n\n"
            "【答案】B\n【难度】易\n【解析】√2 是无理数。\n"
        )
        files = write_exam_files(user, conversation_id, student, answer)
        self.assertEqual(len(files), 2)
        client = TestClient(app)
        logged = client.post(
            "/login",
            data={"username": "teacher", "password": "test-password-1"},
            follow_redirects=False,
        )
        self.assertEqual(logged.status_code, 303)
        paper = client.get(files[0]["url"])
        answer_file = client.get(files[1]["url"])
        self.assertEqual(paper.status_code, 200)
        self.assertEqual(answer_file.status_code, 200)
        paper_text = "\n".join(p.text for p in Document(BytesIO(paper.content)).paragraphs)
        answer_text = "\n".join(p.text for p in Document(BytesIO(answer_file.content)).paragraphs)
        self.assertIn("无理数", paper_text)
        self.assertNotIn("【答案】", paper_text)
        self.assertIn("【答案】", answer_text)


if __name__ == "__main__":
    unittest.main()
