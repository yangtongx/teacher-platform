"""把课件大纲填进课件模板，做成一课时的 PPT。

科目样式用 Apollo's Templates 的 Coursebook：封面和「标题加要点」两页，
只替换文字，留下原有的横线、配色和装饰。老师上传的课件也按同样办法填，
不把旧课文抄进来。
"""

from __future__ import annotations

import copy
import io
import secrets
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

from app import db
from app.config import DATA_DIR

ROLES = ("课题与目标", "导入", "新知", "例题", "练习", "小结", "作业")
SOURCES = ("教案", "课本", "要求")
MAX_SLIDES = 16
MAX_BULLETS = 5
MAX_BULLET = 36
MAX_TITLE = 28

def review_slides(slides: list) -> list[str]:
    """核对每一页都能指回教案、课本或老师的要求。新知不能只凭要求来写。"""
    if not isinstance(slides, list) or not slides:
        return ["没有幻灯片"]
    problems = []
    if len(slides) > MAX_SLIDES:
        problems.append(f"一课时不超过 {MAX_SLIDES} 页")
    roles = []
    for index, slide in enumerate(slides, 1):
        if not isinstance(slide, dict):
            problems.append(f"第 {index} 页格式不对")
            continue
        role = str(slide.get("role") or "").strip()
        roles.append(role)
        if role not in ROLES:
            problems.append(f"第 {index} 页的环节只能是{'、'.join(ROLES)}")
        title = " ".join(str(slide.get("title") or "").split())
        if not title:
            problems.append(f"第 {index} 页没有标题")
        elif len(title) > MAX_TITLE:
            problems.append(f"第 {index} 页标题太长")
        bullets = slide.get("bullets")
        if not isinstance(bullets, list) or not [item for item in bullets if str(item).strip()]:
            problems.append(f"第 {index} 页需要 1 到 {MAX_BULLETS} 条短句")
        elif len([item for item in bullets if str(item).strip()]) > MAX_BULLETS:
            problems.append(f"第 {index} 页超过 {MAX_BULLETS} 条")
        else:
            for item in bullets:
                if len(" ".join(str(item).split())) > MAX_BULLET:
                    problems.append(f"第 {index} 页有一条太长")
                    break
        source = str(slide.get("source") or "").strip()
        if source not in SOURCES:
            problems.append(f"第 {index} 页要写明依据是教案、课本或要求")
        elif role == "新知" and source == "要求":
            problems.append(f"第 {index} 页新知只能依据教案或课本")
    if roles and roles[0] != "课题与目标":
        problems.append("第一页应是课题与目标")
    if "小结" not in roles or "作业" not in roles:
        problems.append("课件需要有小结和作业")
    return problems


def write_ppt(
    user,
    conversation_id: int,
    title: str,
    slides: list,
    template: Path | None = None,
    subject: str = "通用",
) -> dict:
    """核对通过后按选定的课件样式写成 PPT，挂到当前对话。"""
    owned = db.conversation_owner(conversation_id)
    if owned is None or int(owned["user_id"]) != int(user["id"]):
        raise ValueError("这份对话不能导出课件。")
    problems = review_slides(slides)
    if problems:
        raise ValueError("；".join(problems))
    message_id = db.latest_message_id(conversation_id)
    if message_id is None:
        raise ValueError("找不到这次对话，课件没有保存。")
    name = _safe_name(title or str(slides[0].get("title") or "课件"))
    token = secrets.token_urlsafe(16)
    filename = f"{name}_课件.pptx"
    path = DATA_DIR / "docs" / f"{token}.pptx"
    _render(name, slides, path, template, subject)
    db.save_document(int(user["id"]), message_id, token, filename)
    return {"name": filename, "url": f"/files/{token}"}


def _safe_name(text: str) -> str:
    cleaned = "".join(ch for ch in " ".join(str(text).split()) if ch not in '\\/:*?"<>|')
    return cleaned[:40] or "课件"


def _render(course: str, slides: list, path: Path, template: Path | None, subject: str) -> None:
    """打开模板，复制封面页和正文页，再删掉模板自带的示例页。"""
    from app.ppt_style import template_file

    source = template or template_file(subject)
    if not source.is_file():
        raise ValueError("服务器上还没有这个科目的课件样式。")
    presentation = _open_template(source)
    cover_index, body_index = _prototypes(presentation, source)
    original = len(presentation.slides)
    for index, slide in enumerate(slides):
        page = _duplicate(presentation, cover_index if index == 0 else body_index)
        title = " ".join(str(slide.get("title") or "").split())
        bullets = [" ".join(str(item).split()) for item in slide.get("bullets") or [] if str(item).strip()]
        role = str(slide.get("role") or "")
        if index == 0:
            _fill_cover(page, title, bullets, subject)
        else:
            _fill_body(page, title, bullets, role)
        notes = " ".join(str(slide.get("notes") or "").split())
        source_name = str(slide.get("source") or "")
        note_text = "\n".join(part for part in (notes, f"依据：{source_name}", f"课题：{course}") if part)
        page.notes_slide.notes_text_frame.text = note_text
    for _ in range(original):
        _drop_slide(presentation, 0)
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)


def _open_template(path: Path) -> Presentation:
    try:
        return Presentation(str(path))
    except ValueError:
        data = path.read_bytes()
        archive = zipfile.ZipFile(io.BytesIO(data))
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as packed:
            for info in archive.infolist():
                blob = archive.read(info.filename)
                if info.filename == "[Content_Types].xml":
                    blob = blob.replace(
                        b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
                        b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                    )
                packed.writestr(info, blob)
        out.seek(0)
        return Presentation(out)


def _prototypes(presentation: Presentation, path: Path) -> tuple[int, int]:
    """Coursebook / Pencil Box 的第 1 页是封面，第 3 页是标题加要点。"""
    count = len(presentation.slides)
    if count == 0:
        raise ValueError("这份课件样式是空的。")
    name = path.name
    if name.startswith("powerpoint-coursebook") or name.startswith("powerpoint-pencilbox"):
        return 0, 2 if count > 2 else 0
    cover = 0
    best_index = 0
    best_count = -1
    for index, slide in enumerate(presentation.slides):
        if index == 0:
            continue
        boxes = len(_text_boxes(slide))
        if boxes > best_count:
            best_count = boxes
            best_index = index
    return cover, best_index


def _duplicate(presentation: Presentation, index: int):
    source = presentation.slides[index]
    dest = presentation.slides.add_slide(source.slide_layout)
    for shape in list(dest.shapes):
        shape._element.getparent().remove(shape._element)
    mapping: dict[str, str] = {}
    for rel in source.part.rels.values():
        if rel.is_external or "notesSlide" in rel.reltype or "slideLayout" in rel.reltype:
            continue
        mapping[rel.rId] = dest.part.rels.get_or_add(rel.reltype, rel.target_part)
    tree = dest.shapes._spTree
    for shape in source.shapes:
        cloned = copy.deepcopy(shape._element)
        _remap_refs(cloned, mapping)
        try:
            tree.insert_element_before(cloned, "p:extLst")
        except Exception:
            tree.append(cloned)
    return dest


def _remap_refs(element, mapping: dict[str, str]) -> None:
    """复制过来的图片仍指向原页的关系，改到新页上。"""
    for node in element.iter():
        for attr, value in list(node.attrib.items()):
            if value not in mapping:
                continue
            if attr.endswith("}embed") or attr.endswith("}link"):
                node.attrib[attr] = mapping[value]


def _drop_slide(presentation: Presentation, index: int) -> None:
    slide_id = presentation.slides._sldIdLst[index]
    relation = slide_id.get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    )
    presentation.part.drop_rel(relation)
    presentation.slides._sldIdLst.remove(slide_id)


def _fill_cover(page, title: str, bullets: list[str], subject: str) -> None:
    title_box, body, labels = _groups(page)
    if title_box is not None:
        _write_box(title_box, title)
    if body:
        _write_box(body[0], "\n".join(bullets))
        for box in body[1:]:
            _write_box(box, "")
    label = subject if subject and subject != "通用" else "课堂"
    words = ["课件", label]
    for box, word in zip(labels, words):
        _write_box(box, word)
    for box in labels[len(words) :]:
        _write_box(box, "")


def _fill_body(page, title: str, bullets: list[str], role: str) -> None:
    title_box, body, labels = _groups(page)
    if title_box is not None:
        _write_box(title_box, title)
    for index, box in enumerate(body):
        if index < len(bullets):
            _write_box(box, bullets[index])
        elif index == len(body) - 1 and len(bullets) > len(body):
            _write_box(box, "\n".join(bullets[index:]))
        else:
            _write_box(box, "")
    if body and len(bullets) > len(body):
        _write_box(body[-1], "\n".join(bullets[len(body) - 1 :]))
    if labels:
        _write_box(labels[0], role)
        for box in labels[1:]:
            _write_box(box, "")


def _groups(page):
    """按字号分成标题、正文和页眉小字。没有字号时按从上到下。"""
    boxes = _text_boxes(page)
    if not boxes:
        return None, [], []
    sized = [box for box in boxes if _font_size(box) > 0]
    if not sized:
        ordered = sorted(boxes, key=lambda box: (int(box.top), int(box.left)))
        return ordered[0], ordered[1:], []
    title_box = min(sized, key=lambda box: (-_font_size(box), int(box.top)))
    title_size = _font_size(title_box)
    rest = [box for box in sized if box is not title_box]
    # 封面副标题大约 22 磅，正文要点大约 16 磅以上，再小的是页眉页脚。
    if title_size >= 600000:
        cutoff = 200000
    else:
        cutoff = 180000
    body = sorted((box for box in rest if _font_size(box) >= cutoff), key=lambda box: int(box.top))
    labels = sorted((box for box in rest if _font_size(box) < cutoff), key=lambda box: int(box.top))
    if not body and rest:
        body = sorted(rest, key=lambda box: int(box.top))[:1]
        labels = [box for box in rest if box not in body]
    return title_box, body, labels


def _text_boxes(page) -> list:
    found = []
    for shape in page.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if not shape.text_frame.text.strip() and _font_size(shape) <= 0:
            continue
        found.append(shape)
    return found


def _font_size(shape) -> int:
    if not shape.has_text_frame:
        return 0
    best = 0
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size:
                best = max(best, int(run.font.size))
    return best


def _write_box(shape, text: str) -> None:
    frame = shape.text_frame
    original = frame.text or ""
    prefix = ""
    if text:
        stripped = original.lstrip()
        if stripped[:1] in {"✓", "✎", "✔"}:
            prefix = stripped[:1] + "  "
    value = prefix + text
    paragraphs = list(frame.paragraphs)
    first = paragraphs[0]
    if first.runs:
        first.runs[0].text = value
        _use_cjk(first.runs[0])
        for run in first.runs[1:]:
            run.text = ""
    else:
        first.text = value
        if first.runs:
            _use_cjk(first.runs[0])
    for paragraph in paragraphs[1:]:
        for run in paragraph.runs:
            run.text = ""


def _use_cjk(run) -> None:
    """中文用微软雅黑。模板原字体往往没有汉字。"""
    run.font.name = "微软雅黑"
    r_pr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        node = r_pr.find(qn(f"a:{tag}"))
        if node is None:
            node = r_pr.makeelement(qn(f"a:{tag}"))
            r_pr.append(node)
        node.set("typeface", "微软雅黑")
